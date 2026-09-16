import json
import openpyxl
from collections import defaultdict
from datetime import timedelta, date
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Sum, Q
from django.utils import timezone
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from openpyxl.styles import Font, PatternFill

from ..models import (
    Employee, Department, Position, Skill, EmployeeSkill, MachineAuthorization,
    Shift, Attendance, LeaveType, LeaveRequest, Payslip, WorkSchedule,
    WorkIncident, MedicalVisit, ProtectiveEquipment, ShiftAssignment
)
from ..forms import (
    EmployeeForm, EmployeeDocumentForm, SkillForm, EmployeeSkillForm, MachineAuthorizationForm,
    AttendanceForm, AttendanceBulkForm, LeaveRequestForm, PayslipForm, ShiftAssignmentForm,
    WorkIncidentForm, MedicalVisitForm, ProtectiveEquipmentForm, DepartmentForm, PositionForm, ShiftForm
)

@login_required
def drh_dashboard(request):
    today = timezone.now().date()
    current_month = today.month
    current_year = today.year

    total_employes = Employee.objects.filter(statut='ACTIF').count()
    employes_en_conge = Employee.objects.filter(statut='CONGE').count()
    nouveaux_embauches = Employee.objects.filter(
        date_embauche__month=current_month,
        date_embauche__year=current_year
    ).count()

    date_limite = today + timedelta(days=30)
    contrats_expirants = Employee.objects.filter(
        type_contrat__in=['CDD', 'INTERIM'],
        date_fin_contrat__lte=date_limite,
        date_fin_contrat__gte=today,
        statut='ACTIF'
    ).count()

    pointages_jour = Attendance.objects.filter(date=today)
    presents = pointages_jour.filter(statut='PRESENT').count()
    absents = pointages_jour.filter(statut='ABSENT').count()
    retards = pointages_jour.filter(statut='RETARD').count()

    conges_en_attente = LeaveRequest.objects.filter(
        statut__in=['SOUMISE', 'VALIDEE_N1']
    ).count()

    bulletins_mois = Payslip.objects.filter(mois=current_month, annee=current_year).count()
    bulletins_valides = Payslip.objects.filter(
        mois=current_month, annee=current_year, statut='VALIDE'
    ).count()

    masse_salariale = Payslip.objects.filter(
        mois=current_month, annee=current_year,
        statut__in=['VALIDE', 'PAYE']
    ).aggregate(total=Sum('salaire_net'))['total'] or 0

    incidents_mois = WorkIncident.objects.filter(
        date_incident__month=current_month,
        date_incident__year=current_year
    ).count()

    visites_a_planifier = MedicalVisit.objects.filter(
        date_prochaine_visite__lte=date_limite,
        date_prochaine_visite__gte=today
    ).count()

    repartition_dept = Department.objects.annotate(
        nb_emp=Count('employees', filter=Q(employees__statut='ACTIF'))
    ).filter(nb_emp__gt=0).values('name', 'nb_emp')

    repartition_contrat = Employee.objects.filter(
        statut='ACTIF'
    ).values('type_contrat').annotate(count=Count('id'))

    evolution_effectifs = []
    for i in range(5, -1, -1):
        date_ref = today - timedelta(days=30 * i)
        mois_label = date_ref.strftime('%b %Y')
        count = Employee.objects.filter(
            date_embauche__lte=date_ref, statut='ACTIF'
        ).exclude(date_depart__lt=date_ref).count()
        evolution_effectifs.append({'mois': mois_label, 'count': count})

    employes_recents = Employee.objects.filter(statut='ACTIF').order_by('-date_embauche')[:5]

    conges_recents = LeaveRequest.objects.select_related(
        'employee', 'type_conge'
    ).order_by('-date_demande')[:5]

    context = {
        'total_employes': total_employes,
        'employes_en_conge': employes_en_conge,
        'nouveaux_embauches': nouveaux_embauches,
        'contrats_expirants': contrats_expirants,
        'presents': presents,
        'absents': absents,
        'retards': retards,
        'taux_presence': round(presents / max(total_employes, 1) * 100, 1),
        'conges_en_attente': conges_en_attente,
        'bulletins_mois': bulletins_mois,
        'bulletins_valides': bulletins_valides,
        'masse_salariale': masse_salariale,
        'incidents_mois': incidents_mois,
        'visites_a_planifier': visites_a_planifier,
        'repartition_dept_labels': json.dumps([d['name'] for d in repartition_dept]),
        'repartition_dept_data': json.dumps([d['nb_emp'] for d in repartition_dept]),
        'repartition_contrat': list(repartition_contrat),
        'evolution_mois': json.dumps([e['mois'] for e in evolution_effectifs]),
        'evolution_data': json.dumps([e['count'] for e in evolution_effectifs]),
        'employes_recents': employes_recents,
        'conges_recents': conges_recents,
        'today': today,
    }

    return render(request, 'drh/dashboard.html', context)


@login_required
def employee_list(request):
    search = request.GET.get('q', '')
    department = request.GET.get('department', '')
    position = request.GET.get('position', '')
    statut = request.GET.get('statut', '')
    contrat = request.GET.get('contrat', '')

    employees = Employee.objects.select_related(
        'department', 'position', 'machine_affectee'
    ).order_by('nom', 'prenom')

    if search:
        employees = employees.filter(
            Q(nom__icontains=search) | Q(prenom__icontains=search) |
            Q(matricule__icontains=search) | Q(cin__icontains=search)
        )
    if department:
        employees = employees.filter(department_id=department)
    if position:
        employees = employees.filter(position_id=position)
    if statut:
        employees = employees.filter(statut=statut)
    if contrat:
        employees = employees.filter(type_contrat=contrat)

    total = employees.count()
    actifs = employees.filter(statut='ACTIF').count()

    context = {
        'employees': employees, 'total': total, 'actifs': actifs,
        'departments': Department.objects.filter(is_active=True),
        'positions': Position.objects.filter(is_active=True),
        'statut_choices': Employee.STATUT_CHOICES,
        'contrat_choices': Employee.CONTRAT_CHOICES,
        'search': search,
        'selected_department': department,
        'selected_position': position,
        'selected_statut': statut,
        'selected_contrat': contrat,
    }

    return render(request, 'drh/employee_list.html', context)


@login_required
def employee_detail(request, emp_id):
    employee = get_object_or_404(
        Employee.objects.select_related(
            'department', 'position', 'superieur', 'machine_affectee', 'user'
        ), id=emp_id
    )

    documents = employee.documents.all().order_by('-date_upload')
    competences = employee.competences.select_related('skill').order_by('-level')
    autorisations = employee.autorisations_machine.select_related('machine').all()
    pointages = employee.pointages.order_by('-date')[:30]
    conges = employee.demandes_conge.select_related('type_conge').order_by('-date_demande')[:10]
    solde_conge = employee.solde_conge
    bulletins = employee.bulletins_paie.order_by('-annee', '-mois')[:12]
    visites = employee.visites_medicales.order_by('-date_visite')[:5]
    incidents = employee.incidents.order_by('-date_incident')[:5]
    epi = employee.equipements_protection.order_by('-date_attribution')
    subordonnes = Employee.objects.filter(superieur=employee, statut='ACTIF')

    today = timezone.now().date()
    pointages_mois = employee.pointages.filter(
        date__month=today.month, date__year=today.year
    )
    jours_presents = pointages_mois.filter(statut='PRESENT').count()
    jours_absents = pointages_mois.filter(statut='ABSENT').count()
    total_hs = pointages_mois.aggregate(hs=Sum('heures_supplementaires'))['hs'] or 0

    context = {
        'employee': employee, 'documents': documents,
        'competences': competences, 'autorisations': autorisations,
        'pointages': pointages, 'conges': conges, 'solde_conge': solde_conge,
        'bulletins': bulletins, 'visites': visites, 'incidents': incidents,
        'epi': epi, 'subordonnes': subordonnes,
        'jours_presents': jours_presents, 'jours_absents': jours_absents,
        'total_hs': total_hs,
    }

    return render(request, 'drh/employee_detail.html', context)


@login_required
def employee_create(request):
    if request.method == 'POST':
        form = EmployeeForm(request.POST, request.FILES)
        if form.is_valid():
            employee = form.save()
            messages.success(request, f"Employé {employee.nom_complet} créé avec succès ! Matricule : {employee.matricule}")
            return redirect('employee_detail', emp_id=employee.id)
        else:
            messages.error(request, "Erreur dans le formulaire. Vérifiez les champs.")
    else:
        form = EmployeeForm()

    context = {'form': form, 'titre': 'Nouvel Employé'}
    return render(request, 'drh/employee_form.html', context)


@login_required
def employee_edit(request, emp_id):
    employee = get_object_or_404(Employee, id=emp_id)

    if request.method == 'POST':
        form = EmployeeForm(request.POST, request.FILES, instance=employee)
        if form.is_valid():
            form.save()
            messages.success(request, f"Employé {employee.nom_complet} mis à jour !")
            return redirect('employee_detail', emp_id=employee.id)
    else:
        form = EmployeeForm(instance=employee)

    context = {
        'form': form, 'employee': employee,
        'titre': f'Modifier : {employee.nom_complet}',
    }
    return render(request, 'drh/employee_form.html', context)


@login_required
def employee_document_add(request, emp_id):
    employee = get_object_or_404(Employee, id=emp_id)

    if request.method == 'POST':
        form = EmployeeDocumentForm(request.POST, request.FILES)
        if form.is_valid():
            doc = form.save(commit=False)
            doc.employee = employee
            doc.save()
            messages.success(request, f"Document '{doc.nom}' ajouté.")
            return redirect('employee_detail', emp_id=emp_id)
    else:
        form = EmployeeDocumentForm()

    context = {'form': form, 'employee': employee, 'titre': 'Ajouter un document'}
    return render(request, 'drh/document_form.html', context)


@login_required
def skill_list(request):
    skills = Skill.objects.select_related('machine_associee').order_by('category', 'name')

    if request.method == 'POST':
        form = SkillForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Compétence créée !")
            return redirect('skill_list')
    else:
        form = SkillForm()

    context = {'skills': skills, 'form': form}
    return render(request, 'drh/skill_list.html', context)


@login_required
def employee_skill_add(request, emp_id):
    employee = get_object_or_404(Employee, id=emp_id)

    if request.method == 'POST':
        form = EmployeeSkillForm(request.POST, request.FILES)
        if form.is_valid():
            emp_skill = form.save(commit=False)
            emp_skill.employee = employee
            emp_skill.save()
            messages.success(request, f"Compétence '{emp_skill.skill.name}' ajoutée.")
            return redirect('employee_detail', emp_id=emp_id)
    else:
        form = EmployeeSkillForm()

    context = {'form': form, 'employee': employee, 'titre': 'Ajouter une compétence'}
    return render(request, 'drh/skill_form.html', context)


@login_required
def machine_authorization_add(request, emp_id):
    employee = get_object_or_404(Employee, id=emp_id)

    if request.method == 'POST':
        form = MachineAuthorizationForm(request.POST)
        if form.is_valid():
            auth = form.save(commit=False)
            auth.employee = employee
            auth.save()
            messages.success(request, f"Autorisation pour {auth.machine.name} créée.")
            return redirect('employee_detail', emp_id=emp_id)
    else:
        form = MachineAuthorizationForm()

    context = {'form': form, 'employee': employee, 'titre': 'Autorisation machine'}
    return render(request, 'drh/authorization_form.html', context)


@login_required
def machine_authorization_validate(request, auth_id):
    auth = get_object_or_404(MachineAuthorization, id=auth_id)

    if request.method == 'POST':
        try:
            validateur = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            validateur = None

        auth.statut = 'VALIDE'
        auth.date_validation = timezone.now().date()
        auth.validateur = validateur
        auth.save()
        messages.success(request, f"Autorisation validée pour {auth.employee.nom_complet}")

    return redirect('employee_detail', emp_id=auth.employee.id)


@login_required
def attendance_list(request):
    date_filter = request.GET.get('date', '')
    department = request.GET.get('department', '')
    shift = request.GET.get('shift', '')
    statut = request.GET.get('statut', '')

    if not date_filter:
        date_filter = timezone.now().date().isoformat()

    attendances = Attendance.objects.select_related(
        'employee', 'employee__department', 'shift', 'machine'
    ).order_by('-date', 'employee__nom')

    if date_filter:
        attendances = attendances.filter(date=date_filter)
    if department:
        attendances = attendances.filter(employee__department_id=department)
    if shift:
        attendances = attendances.filter(shift_id=shift)
    if statut:
        attendances = attendances.filter(statut=statut)

    stats = {
        'total': attendances.count(),
        'presents': attendances.filter(statut='PRESENT').count(),
        'absents': attendances.filter(statut='ABSENT').count(),
        'retards': attendances.filter(statut='RETARD').count(),
        'conges': attendances.filter(statut='CONGE').count(),
    }

    context = {
        'attendances': attendances[:200], 'stats': stats,
        'date_filter': date_filter,
        'departments': Department.objects.filter(is_active=True),
        'shifts': Shift.objects.filter(is_active=True),
        'statut_choices': Attendance.STATUT_CHOICES,
        'selected_department': department,
        'selected_shift': shift,
        'selected_statut': statut,
    }

    return render(request, 'drh/attendance_list.html', context)


@login_required
def attendance_create(request):
    if request.method == 'POST':
        form = AttendanceForm(request.POST)
        if form.is_valid():
            attendance = form.save()
            attendance.calculer_heures()
            messages.success(request, "Pointage enregistré !")
            return redirect('attendance_list')
    else:
        form = AttendanceForm(initial={'date': timezone.now().date()})

    context = {'form': form, 'titre': 'Nouveau pointage'}
    return render(request, 'drh/attendance_form.html', context)


@login_required
def attendance_bulk(request):
    if request.method == 'POST':
        date = request.POST.get('date')
        shift_id = request.POST.get('shift')
        department_id = request.POST.get('department')

        employees = Employee.objects.filter(statut='ACTIF')
        if department_id:
            employees = employees.filter(department_id=department_id)

        shift = Shift.objects.get(id=shift_id) if shift_id else None
        created = 0

        for emp in employees:
            if not Attendance.objects.filter(employee=emp, date=date).exists():
                statut = request.POST.get(f'statut_{emp.id}', 'PRESENT')
                heure_arrivee = request.POST.get(f'arrivee_{emp.id}') or None
                heure_depart = request.POST.get(f'depart_{emp.id}') or None

                attendance = Attendance.objects.create(
                    employee=emp, date=date, shift=shift,
                    statut=statut, heure_arrivee=heure_arrivee,
                    heure_depart=heure_depart,
                    machine=emp.machine_affectee,
                )
                attendance.calculer_heures()
                created += 1

        messages.success(request, f"{created} pointages créés !")
        return redirect('attendance_list')

    form = AttendanceBulkForm()
    employees = Employee.objects.filter(statut='ACTIF').select_related('department', 'position')

    context = {'form': form, 'employees': employees, 'titre': 'Pointage en masse'}
    return render(request, 'drh/attendance_bulk.html', context)


@login_required
def leave_list(request):
    statut = request.GET.get('statut', '')
    type_conge = request.GET.get('type', '')
    department = request.GET.get('department', '')

    leaves = LeaveRequest.objects.select_related(
        'employee', 'employee__department', 'type_conge',
        'validateur_n1', 'validateur_rh'
    ).order_by('-date_demande')

    if statut:
        leaves = leaves.filter(statut=statut)
    if type_conge:
        leaves = leaves.filter(type_conge_id=type_conge)
    if department:
        leaves = leaves.filter(employee__department_id=department)

    stats = {
        'en_attente': leaves.filter(statut__in=['SOUMISE', 'VALIDEE_N1']).count(),
        'validees': leaves.filter(statut='VALIDEE_RH').count(),
        'refusees': leaves.filter(statut='REFUSEE').count(),
    }

    context = {
        'leaves': leaves[:100], 'stats': stats,
        'statut_choices': LeaveRequest.STATUT_CHOICES,
        'types': LeaveType.objects.filter(is_active=True),
        'departments': Department.objects.filter(is_active=True),
        'selected_statut': statut,
        'selected_type': type_conge,
        'selected_department': department,
    }

    return render(request, 'drh/leave_list.html', context)


@login_required
def leave_create(request, emp_id=None):
    employee = None
    if emp_id:
        employee = get_object_or_404(Employee, id=emp_id)

    if request.method == 'POST':
        form = LeaveRequestForm(request.POST, request.FILES)
        emp_id_form = request.POST.get('employee')

        if form.is_valid():
            leave = form.save(commit=False)
            if employee:
                leave.employee = employee
            else:
                leave.employee = get_object_or_404(Employee, id=emp_id_form)
            leave.statut = 'SOUMISE'
            leave.save()
            messages.success(request, "Demande de congé soumise !")
            return redirect('leave_list')
    else:
        form = LeaveRequestForm()

    context = {
        'form': form, 'employee': employee,
        'employees': Employee.objects.filter(statut='ACTIF') if not employee else None,
        'titre': 'Nouvelle demande de congé',
    }
    return render(request, 'drh/leave_form.html', context)


@login_required
def leave_validate_n1(request, leave_id):
    leave = get_object_or_404(LeaveRequest, id=leave_id)

    if request.method == 'POST':
        try:
            validateur = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            validateur = None

        leave.valider_n1(validateur)
        messages.success(request, f"Congé validé (N+1) pour {leave.employee.nom_complet}")

    return redirect('leave_list')


@login_required
def leave_validate_rh(request, leave_id):
    leave = get_object_or_404(LeaveRequest, id=leave_id)

    if request.method == 'POST':
        try:
            validateur = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            validateur = None

        leave.valider_rh(validateur)
        messages.success(request, f"Congé validé (RH) pour {leave.employee.nom_complet}. Solde déduit : {leave.nb_jours} jours")

    return redirect('leave_list')


@login_required
def leave_reject(request, leave_id):
    leave = get_object_or_404(LeaveRequest, id=leave_id)

    if request.method == 'POST':
        motif = request.POST.get('motif_refus', '')
        try:
            validateur = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            validateur = None

        leave.refuser(validateur, motif)
        messages.warning(request, f"Congé refusé pour {leave.employee.nom_complet}")

    return redirect('leave_list')


@login_required
def payslip_list(request):
    mois = request.GET.get('mois', timezone.now().month)
    annee = request.GET.get('annee', timezone.now().year)
    department = request.GET.get('department', '')
    statut = request.GET.get('statut', '')

    payslips = Payslip.objects.select_related(
        'employee', 'employee__department', 'employee__position'
    ).filter(mois=mois, annee=annee).order_by('employee__nom')

    if department:
        payslips = payslips.filter(employee__department_id=department)
    if statut:
        payslips = payslips.filter(statut=statut)

    totaux = payslips.aggregate(
        total_brut=Sum('salaire_brut'),
        total_net=Sum('salaire_net'),
        total_cotisations=Sum('total_cotisations'),
        total_irg=Sum('irg'),
        total_charges=Sum('total_charges_patronales'),
    )

    context = {
        'payslips': payslips, 'totaux': totaux,
        'mois': int(mois), 'annee': int(annee),
        'mois_list': [(i, f"{i:02d}") for i in range(1, 13)],
        'annee_list': range(timezone.now().year - 2, timezone.now().year + 1),
        'departments': Department.objects.filter(is_active=True),
        'statut_choices': Payslip.STATUT_CHOICES,
        'selected_department': department,
        'selected_statut': statut,
    }

    return render(request, 'drh/payslip_list.html', context)


@login_required
def payslip_create(request):
    if request.method == 'POST':
        form = PayslipForm(request.POST)
        if form.is_valid():
            payslip = form.save()
            payslip.calculer()
            messages.success(request, f"Bulletin créé et calculé pour {payslip.employee.nom_complet}")
            return redirect('payslip_detail', slip_id=payslip.id)
    else:
        form = PayslipForm(initial={
            'mois': timezone.now().month,
            'annee': timezone.now().year,
            'jours_travailles': 26,
        })

    context = {'form': form, 'titre': 'Nouveau bulletin de paie'}
    return render(request, 'drh/payslip_form.html', context)


@login_required
def payslip_detail(request, slip_id):
    payslip = get_object_or_404(
        Payslip.objects.select_related('employee', 'employee__department', 'employee__position'),
        id=slip_id
    )
    context = {'payslip': payslip}
    return render(request, 'drh/payslip_detail.html', context)


@login_required
def payslip_calculate(request, slip_id):
    payslip = get_object_or_404(Payslip, id=slip_id)
    if request.method == 'POST':
        payslip.calculer()
        messages.success(request, "Bulletin recalculé !")
    return redirect('payslip_detail', slip_id=slip_id)


@login_required
def payslip_validate(request, slip_id):
    payslip = get_object_or_404(Payslip, id=slip_id)

    if request.method == 'POST':
        try:
            validateur = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            validateur = None

        payslip.statut = 'VALIDE'
        payslip.date_validation = timezone.now()
        payslip.valide_par = validateur
        payslip.save()
        messages.success(request, "Bulletin validé !")

    return redirect('payslip_detail', slip_id=slip_id)


@login_required
def payslip_bulk_generate(request):
    if request.method == 'POST':
        mois = int(request.POST.get('mois', timezone.now().month))
        annee = int(request.POST.get('annee', timezone.now().year))
        department_id = request.POST.get('department')

        employees = Employee.objects.filter(statut='ACTIF')
        if department_id:
            employees = employees.filter(department_id=department_id)

        created = 0
        for emp in employees:
            if not Payslip.objects.filter(employee=emp, mois=mois, annee=annee).exists():
                pointages = Attendance.objects.filter(
                    employee=emp, date__month=mois, date__year=annee
                )
                jours_presents = pointages.filter(statut='PRESENT').count()
                jours_absents = pointages.filter(statut='ABSENT').count()
                heures_sup = pointages.aggregate(hs=Sum('heures_supplementaires'))['hs'] or 0
                heures_nuit = pointages.aggregate(hn=Sum('heures_nuit'))['hn'] or 0

                payslip = Payslip.objects.create(
                    employee=emp, mois=mois, annee=annee,
                    jours_travailles=jours_presents,
                    jours_absence=jours_absents,
                    heures_supplementaires_25=heures_sup,
                    heures_nuit=heures_nuit,
                )
                payslip.calculer()
                created += 1

        messages.success(request, f"{created} bulletins générés pour {mois:02d}/{annee} !")
        return redirect('payslip_list')

    context = {
        'mois': timezone.now().month,
        'annee': timezone.now().year,
        'mois_list': [(i, f"{i:02d}") for i in range(1, 13)],
        'departments': Department.objects.filter(is_active=True),
    }
    return render(request, 'drh/payslip_bulk.html', context)


@login_required
def schedule_list(request):
    schedules = WorkSchedule.objects.select_related(
        'department', 'machine', 'cree_par'
    ).order_by('-date_debut')
    context = {'schedules': schedules}
    return render(request, 'drh/schedule_list.html', context)


@login_required
def schedule_detail(request, schedule_id):
    schedule = get_object_or_404(WorkSchedule, id=schedule_id)

    affectations = schedule.affectations.select_related(
        'employee', 'shift', 'machine'
    ).order_by('date', 'shift__heure_debut')

    affectations_par_date = defaultdict(list)
    for aff in affectations:
        affectations_par_date[aff.date].append(aff)

    context = {
        'schedule': schedule,
        'affectations_par_date': dict(affectations_par_date),
        'shifts': Shift.objects.filter(is_active=True),
        'employees': Employee.objects.filter(statut='ACTIF').order_by('nom'),
    }
    return render(request, 'drh/schedule_detail.html', context)


@login_required
def shift_assignment_add(request, schedule_id):
    schedule = get_object_or_404(WorkSchedule, id=schedule_id)

    if request.method == 'POST':
        form = ShiftAssignmentForm(request.POST)
        if form.is_valid():
            assignment = form.save(commit=False)
            assignment.schedule = schedule
            assignment.save()
            messages.success(request, "Affectation ajoutée !")
            return redirect('schedule_detail', schedule_id=schedule_id)
    else:
        form = ShiftAssignmentForm()

    context = {'form': form, 'schedule': schedule, 'titre': 'Nouvelle affectation'}
    return render(request, 'drh/assignment_form.html', context)


@login_required
def incident_list(request):
    type_incident = request.GET.get('type', '')
    gravite = request.GET.get('gravite', '')
    cloture = request.GET.get('cloture', '')

    incidents = WorkIncident.objects.select_related(
        'employee', 'machine', 'cree_par'
    ).order_by('-date_incident')

    if type_incident:
        incidents = incidents.filter(type_incident=type_incident)
    if gravite:
        incidents = incidents.filter(gravite=gravite)
    if cloture:
        incidents = incidents.filter(cloture=(cloture == '1'))

    stats = {
        'total': incidents.count(),
        'accidents': incidents.filter(type_incident='ACCIDENT').count(),
        'non_clotures': incidents.filter(cloture=False).count(),
        'jours_arret': incidents.aggregate(j=Sum('jours_arret'))['j'] or 0,
    }

    context = {
        'incidents': incidents[:100], 'stats': stats,
        'type_choices': WorkIncident.TYPE_CHOICES,
        'gravite_choices': WorkIncident.GRAVITE_CHOICES,
        'selected_type': type_incident,
        'selected_gravite': gravite,
        'selected_cloture': cloture,
    }
    return render(request, 'drh/incident_list.html', context)


@login_required
def incident_create(request):
    if request.method == 'POST':
        form = WorkIncidentForm(request.POST, request.FILES)
        if form.is_valid():
            incident = form.save(commit=False)
            try:
                incident.cree_par = Employee.objects.get(user=request.user)
            except Employee.DoesNotExist:
                pass
            incident.save()
            messages.success(request, f"Incident {incident.reference} déclaré !")
            return redirect('incident_list')
    else:
        form = WorkIncidentForm()

    context = {'form': form, 'titre': 'Déclarer un incident'}
    return render(request, 'drh/incident_form.html', context)


@login_required
def incident_detail(request, incident_id):
    incident = get_object_or_404(
        WorkIncident.objects.select_related('employee', 'machine', 'cree_par'),
        id=incident_id
    )
    context = {'incident': incident}
    return render(request, 'drh/incident_detail.html', context)


@login_required
def medical_visit_list(request):
    visits = MedicalVisit.objects.select_related('employee').order_by('-date_visite')
    today = timezone.now().date()
    date_limite = today + timedelta(days=30)
    a_planifier = visits.filter(
        date_prochaine_visite__lte=date_limite,
        date_prochaine_visite__gte=today
    )
    context = {'visits': visits[:100], 'a_planifier': a_planifier}
    return render(request, 'drh/medical_list.html', context)


@login_required
def medical_visit_create(request):
    if request.method == 'POST':
        form = MedicalVisitForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, "Visite médicale enregistrée !")
            return redirect('medical_list')
    else:
        form = MedicalVisitForm()
    context = {'form': form, 'titre': 'Nouvelle visite médicale'}
    return render(request, 'drh/medical_form.html', context)


@login_required
def epi_list(request):
    epis = ProtectiveEquipment.objects.select_related('employee').order_by('-date_attribution')
    today = timezone.now().date()
    expires = epis.filter(date_expiration__lte=today)
    context = {'epis': epis[:100], 'expires': expires}
    return render(request, 'drh/epi_list.html', context)


@login_required
def epi_create(request):
    if request.method == 'POST':
        form = ProtectiveEquipmentForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "EPI attribué !")
            return redirect('epi_list')
    else:
        form = ProtectiveEquipmentForm()
    context = {'form': form, 'titre': 'Attribuer un EPI'}
    return render(request, 'drh/epi_form.html', context)


@login_required
def department_list(request):
    departments = Department.objects.annotate(
        nb_emp=Count('employees', filter=Q(employees__statut='ACTIF'))
    ).order_by('name')

    if request.method == 'POST':
        form = DepartmentForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Département créé !")
            return redirect('department_list')
    else:
        form = DepartmentForm()

    context = {'departments': departments, 'form': form}
    return render(request, 'drh/department_list.html', context)


@login_required
def position_list(request):
    positions = Position.objects.select_related('department').annotate(
        nb_emp=Count('employees', filter=Q(employees__statut='ACTIF'))
    ).order_by('category', 'name')

    if request.method == 'POST':
        form = PositionForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Poste créé !")
            return redirect('position_list')
    else:
        form = PositionForm()

    context = {'positions': positions, 'form': form}
    return render(request, 'drh/position_list.html', context)


@login_required
def shift_list(request):
    shifts = Shift.objects.all().order_by('heure_debut')

    if request.method == 'POST':
        form = ShiftForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Équipe créée !")
            return redirect('shift_list')
    else:
        form = ShiftForm()

    context = {'shifts': shifts, 'form': form}
    return render(request, 'drh/shift_list.html', context)


@login_required
def export_employees_excel(request):
    employees = Employee.objects.select_related(
        'department', 'position'
    ).filter(statut='ACTIF').order_by('nom')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employés"

    headers = ['Matricule', 'Nom', 'Prénom', 'CIN', 'Département', 'Poste',
               'Date embauche', 'Type contrat', 'Salaire base', 'Téléphone', 'Email']

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row, emp in enumerate(employees, 2):
        ws.cell(row=row, column=1, value=emp.matricule)
        ws.cell(row=row, column=2, value=emp.nom)
        ws.cell(row=row, column=3, value=emp.prenom)
        ws.cell(row=row, column=4, value=emp.cin)
        ws.cell(row=row, column=5, value=emp.department.name if emp.department else '')
        ws.cell(row=row, column=6, value=emp.position.name if emp.position else '')
        ws.cell(row=row, column=7, value=emp.date_embauche.strftime('%d/%m/%Y') if emp.date_embauche else '')
        ws.cell(row=row, column=8, value=emp.get_type_contrat_display())
        ws.cell(row=row, column=9, value=float(emp.salaire_base))
        ws.cell(row=row, column=10, value=emp.telephone)
        ws.cell(row=row, column=11, value=emp.email)

    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="employes.xlsx"'
    wb.save(response)
    return response


@login_required
def export_payslips_excel(request):
    mois = request.GET.get('mois', timezone.now().month)
    annee = request.GET.get('annee', timezone.now().year)

    payslips = Payslip.objects.select_related(
        'employee', 'employee__department'
    ).filter(mois=mois, annee=annee).order_by('employee__nom')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Paie {mois:02d}-{annee}"

    headers = ['Matricule', 'Nom Complet', 'Département', 'Salaire Base',
               'Primes', 'Brut', 'Cotisations', 'IRG', 'Retenues', 'Net']

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row, slip in enumerate(payslips, 2):
        primes = float(slip.prime_rendement + slip.prime_presence + slip.prime_nuit +
                       slip.prime_anciennete + slip.prime_transport + slip.prime_panier +
                       slip.heures_sup_montant + slip.autres_primes)

        ws.cell(row=row, column=1, value=slip.employee.matricule)
        ws.cell(row=row, column=2, value=slip.employee.nom_complet)
        ws.cell(row=row, column=3, value=slip.employee.department.name if slip.employee.department else '')
        ws.cell(row=row, column=4, value=float(slip.salaire_base))
        ws.cell(row=row, column=5, value=primes)
        ws.cell(row=row, column=6, value=float(slip.salaire_brut))
        ws.cell(row=row, column=7, value=float(slip.total_cotisations))
        ws.cell(row=row, column=8, value=float(slip.irg))
        ws.cell(row=row, column=9, value=float(slip.total_retenues))
        ws.cell(row=row, column=10, value=float(slip.salaire_net))

    last_row = len(payslips) + 2
    ws.cell(row=last_row, column=1, value="TOTAUX")
    ws.cell(row=last_row, column=1).font = Font(bold=True)
    for col in range(4, 11):
        ws.cell(row=last_row, column=col, value=f"=SUM({openpyxl.utils.get_column_letter(col)}2:{openpyxl.utils.get_column_letter(col)}{last_row-1})")
        ws.cell(row=last_row, column=col).font = Font(bold=True)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="paie_{mois:02d}_{annee}.xlsx"'
    wb.save(response)
    return response