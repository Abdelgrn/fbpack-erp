"""Vues superadministrateur du centre de surveillance des connexions."""

import ipaddress
import json
import logging
from datetime import datetime, time, timedelta
from functools import wraps

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import logout as auth_logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_GET, require_POST

from ..models.permissions import AdministrationAuditLog
from ..models.security_audit import (
    AccountSecurityState,
    KnownAccessSignature,
    SecurityEvent,
    SecurityEventType,
    SecurityRiskLevel,
    SecuritySession,
)
from ..services.security_monitoring import (
    current_session_fingerprint,
    disable_user_temporarily,
    fingerprint_ip,
    get_request_network,
    mark_signature_recognized,
    reconcile_expired_sessions,
    revoke_all_user_sessions,
    revoke_security_session,
)

logger = logging.getLogger("security")


def superadmin_required(view_func):
    """Contrôle serveur strict : aucune permission de module ne suffit."""

    @wraps(view_func)
    def checked(request, *args, **kwargs):
        if not request.user.is_superuser:
            raise PermissionDenied("Centre de sécurité réservé aux superadministrateurs.")
        return view_func(request, *args, **kwargs)

    return login_required(checked)


def _local_day_bounds(day_value, end=False):
    parsed = parse_date(day_value or "")
    if not parsed:
        return None
    local_value = datetime.combine(parsed, time.max if end else time.min)
    return timezone.make_aware(local_value, timezone.get_current_timezone())


def _valid_ip(value):
    value = (value or "").strip()
    if not value:
        return ""
    try:
        return str(ipaddress.ip_address(value))
    except ValueError:
        return ""


def _filtered_events(request):
    events = SecurityEvent.objects.select_related("user", "session")
    event_type = (request.GET.get("event_type") or "").strip()
    risk = (request.GET.get("risk") or "").strip()
    user_id = (request.GET.get("user") or "").strip()
    ip_value = (request.GET.get("ip") or "").strip()
    route = (request.GET.get("route") or "").strip()[:120]
    query = (request.GET.get("q") or "").strip()[:120]

    valid_types = {choice for choice, _ in SecurityEventType.choices}
    valid_risks = {choice for choice, _ in SecurityRiskLevel.choices}
    if event_type in valid_types:
        events = events.filter(event_type=event_type)
    if risk in valid_risks:
        events = events.filter(risk_level=risk)
    if user_id.isdigit():
        events = events.filter(user_id=int(user_id))
    if ip_value:
        normalized_ip = _valid_ip(ip_value)
        events = events.filter(client_ip=normalized_ip) if normalized_ip else events.none()
    if route:
        events = events.filter(route__icontains=route)
    if query:
        events = events.filter(
            Q(user__username__icontains=query)
            | Q(reason__icontains=query)
            | Q(device__icontains=query)
            | Q(browser__icontains=query)
            | Q(operating_system__icontains=query)
            | Q(route__icontains=query)
        )

    date_from = _local_day_bounds(request.GET.get("date_from"), end=False)
    date_to = _local_day_bounds(request.GET.get("date_to"), end=True)
    if date_from:
        events = events.filter(occurred_at__gte=date_from)
    if date_to:
        events = events.filter(occurred_at__lte=date_to)
    return events


def _record_admin_action(request, action, target=None, details=""):
    client_ip, _ = get_request_network(request)
    AdministrationAuditLog.objects.create(
        actor=request.user,
        target_user=target,
        action=action[:40],
        module="security",
        details=(details or "")[:2000],
        ip_address=client_ip,
    )


def _password_and_confirmation_valid(request, expected_confirmation):
    if not request.user.check_password(request.POST.get("current_password", "")):
        messages.error(request, "Mot de passe du superadministrateur incorrect.")
        return False
    supplied = (request.POST.get("confirmation") or "").strip().upper()
    if supplied != expected_confirmation.upper():
        messages.error(
            request,
            f"Confirmation invalide. Saisissez exactement : {expected_confirmation}",
        )
        return False
    return True


def _monitoring_redirect():
    return redirect("security_monitoring")


@superadmin_required
@require_GET
def security_monitoring(request):
    reconcile_expired_sessions()
    now = timezone.now()
    online_minutes = max(1, int(getattr(settings, "SECURITY_MONITOR_ONLINE_MINUTES", 5)))
    online_cutoff = now - timedelta(minutes=online_minutes)
    day_cutoff = now - timedelta(hours=24)

    sessions = list(
        SecuritySession.objects.select_related("user", "revoked_by")
        .filter(
            Q(status=SecuritySession.Status.ACTIVE)
            | Q(ended_at__gte=day_cutoff)
        )
        .order_by("-last_activity_at")[:120]
    )
    current_fingerprint = current_session_fingerprint(request)

    user_ids = {item.user_id for item in sessions if item.user_id}
    signatures = KnownAccessSignature.objects.filter(user_id__in=user_ids)
    signature_map = {
        (item.user_id, item.kind, item.fingerprint): item
        for item in signatures
    }
    states = {
        item.user_id: item
        for item in AccountSecurityState.objects.filter(user_id__in=user_ids)
    }

    for item in sessions:
        if item.status == SecuritySession.Status.ACTIVE:
            item.monitor_status = "ONLINE" if item.last_activity_at >= online_cutoff else "IDLE"
        else:
            item.monitor_status = "OFFLINE"
        item.is_current_session = item.session_fingerprint == current_fingerprint
        item.short_session_fingerprint = item.session_fingerprint[:16]
        item.known_ip = signature_map.get(
            (item.user_id, KnownAccessSignature.Kind.IP, fingerprint_ip(item.client_ip))
        )
        item.known_device = signature_map.get(
            (item.user_id, KnownAccessSignature.Kind.DEVICE, item.device_fingerprint)
        )
        item.account_security_state = states.get(item.user_id)

    events_query = _filtered_events(request)
    paginator = Paginator(events_query, 50)
    event_page = paginator.get_page(request.GET.get("page"))
    preserved = request.GET.copy()
    preserved.pop("page", None)

    context = {
        "page_title": "Surveillance des connexions",
        "sessions": sessions,
        "events_page": event_page,
        "event_choices": SecurityEventType.choices,
        "risk_choices": SecurityRiskLevel.choices,
        "filter_users": User.objects.order_by("username").only(
            "id", "username", "is_superuser", "is_active"
        ),
        "query_without_page": preserved.urlencode(),
        "online_minutes": online_minutes,
        "security_event_retention_days": int(
            getattr(settings, "SECURITY_EVENT_RETENTION_DAYS", 90)
        ),
        "stats": {
            "online": SecuritySession.objects.filter(
                status=SecuritySession.Status.ACTIVE,
                last_activity_at__gte=online_cutoff,
            ).count(),
            "failures_24h": SecurityEvent.objects.filter(
                event_type__in=[
                    SecurityEventType.LOGIN_FAILED,
                    SecurityEventType.LOGIN_LOCKED,
                    SecurityEventType.ACCOUNT_DISABLED_ATTEMPT,
                ],
                occurred_at__gte=day_cutoff,
            ).count(),
            "critical_24h": SecurityEvent.objects.filter(
                risk_level=SecurityRiskLevel.CRITICAL,
                occurred_at__gte=day_cutoff,
            ).count(),
            "new_ips_24h": SecurityEvent.objects.filter(
                event_type=SecurityEventType.NEW_IP,
                occurred_at__gte=day_cutoff,
            ).count(),
        },
    }
    return render(request, "security/security_monitoring.html", context)


@superadmin_required
@require_POST
def revoke_session_view(request, session_id):
    monitored = get_object_or_404(
        SecuritySession.objects.select_related("user"),
        pk=session_id,
    )
    if monitored.status != SecuritySession.Status.ACTIVE:
        messages.info(request, "Cette session est déjà fermée.")
        return _monitoring_redirect()

    is_current = monitored.session_fingerprint == current_session_fingerprint(request)
    expected = "REVOQUER MA SESSION" if is_current else "REVOQUER"
    if not _password_and_confirmation_valid(request, expected):
        return _monitoring_redirect()

    reason = (request.POST.get("reason") or "Révocation depuis le centre de sécurité").strip()[:255]
    revoke_security_session(monitored, request.user, reason)
    _record_admin_action(
        request,
        "SECURITY_SESSION_REVOKED",
        monitored.user,
        f"Session {monitored.session_fingerprint[:16]} — {reason}",
    )
    logger.warning(
        "security_session_revoked actor_id=%s target_user_id=%s session_id=%s",
        request.user.pk,
        monitored.user_id,
        monitored.pk,
    )

    if is_current:
        auth_logout(request)
        messages.warning(request, "Votre session a été révoquée comme demandé.")
        return redirect(f"{settings.LOGIN_URL}?session_revoked=1")

    messages.success(request, "Session révoquée avec succès.")
    return _monitoring_redirect()


@superadmin_required
@require_POST
def revoke_user_sessions_view(request, user_id):
    target = get_object_or_404(User, pk=user_id)
    is_self = target.pk == request.user.pk
    expected = "REVOQUER TOUTES MES SESSIONS" if is_self else "REVOQUER TOUT"
    if not _password_and_confirmation_valid(request, expected):
        return _monitoring_redirect()

    reason = (request.POST.get("reason") or "Révocation globale depuis le centre de sécurité").strip()[:255]
    monitored_count, django_count = revoke_all_user_sessions(
        target,
        request.user,
        reason,
    )
    _record_admin_action(
        request,
        "SECURITY_ALL_SESSIONS_REVOKED",
        target,
        f"{monitored_count} session(s) suivie(s), {django_count} session(s) Django — {reason}",
    )
    logger.warning(
        "security_all_sessions_revoked actor_id=%s target_user_id=%s monitored=%s django=%s",
        request.user.pk,
        target.pk,
        monitored_count,
        django_count,
    )

    if is_self:
        auth_logout(request)
        messages.warning(request, "Toutes vos sessions ont été révoquées comme demandé.")
        return redirect(f"{settings.LOGIN_URL}?session_revoked=1")

    messages.success(
        request,
        f"Toutes les sessions de {target.username} ont été révoquées.",
    )
    return _monitoring_redirect()


@superadmin_required
@require_POST
def disable_user_temporarily_view(request, user_id):
    target = get_object_or_404(User, pk=user_id)
    if target.pk == request.user.pk:
        messages.error(request, "Vous ne pouvez pas désactiver temporairement votre propre compte.")
        return _monitoring_redirect()
    if not target.is_active:
        messages.error(request, "Ce compte est déjà inactif ; aucune réactivation automatique n'a été créée.")
        return _monitoring_redirect()
    if target.is_superuser and User.objects.filter(is_superuser=True, is_active=True).count() <= 1:
        messages.error(request, "Impossible de désactiver le dernier superadministrateur actif.")
        return _monitoring_redirect()

    expected = "DESACTIVER SUPERADMIN" if target.is_superuser else "DESACTIVER"
    if not _password_and_confirmation_valid(request, expected):
        return _monitoring_redirect()

    try:
        minutes = int(request.POST.get("duration_minutes", "60"))
    except (TypeError, ValueError):
        minutes = 60
    minutes = max(5, min(7 * 24 * 60, minutes))
    reason = (request.POST.get("reason") or "").strip()[:255]
    if len(reason) < 3:
        messages.error(request, "Un motif explicite est obligatoire.")
        return _monitoring_redirect()

    state = disable_user_temporarily(
        target,
        request.user,
        minutes,
        reason,
    )
    _record_admin_action(
        request,
        "SECURITY_ACCOUNT_TEMP_DISABLED",
        target,
        f"Jusqu'au {timezone.localtime(state.temporarily_disabled_until):%d/%m/%Y %H:%M} — {reason}",
    )
    logger.warning(
        "security_account_temp_disabled actor_id=%s target_user_id=%s duration_minutes=%s",
        request.user.pk,
        target.pk,
        minutes,
    )
    messages.success(
        request,
        f"Compte {target.username} désactivé temporairement jusqu'au "
        f"{timezone.localtime(state.temporarily_disabled_until):%d/%m/%Y %H:%M}.",
    )
    return _monitoring_redirect()


@superadmin_required
@require_POST
def recognize_signature_view(request, signature_id):
    signature = get_object_or_404(
        KnownAccessSignature.objects.select_related("user"),
        pk=signature_id,
    )
    if not _password_and_confirmation_valid(request, "RECONNAITRE"):
        return _monitoring_redirect()
    if signature.is_recognized:
        messages.info(request, "Cette IP ou cet appareil est déjà reconnu.")
        return _monitoring_redirect()

    mark_signature_recognized(signature, request.user)
    _record_admin_action(
        request,
        "SECURITY_SIGNATURE_RECOGNIZED",
        signature.user,
        f"{signature.get_kind_display()} : {signature.display_value}",
    )
    messages.success(request, f"{signature.get_kind_display()} marqué(e) comme reconnu(e).")
    return _monitoring_redirect()


def _excel_safe(value):
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


@superadmin_required
@require_GET
def export_security_events_excel(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        messages.error(request, "Le paquet openpyxl est requis pour l'export Excel.")
        return _monitoring_redirect()

    limit = max(100, int(getattr(settings, "SECURITY_EVENT_EXPORT_LIMIT", 50000)))
    events = _filtered_events(request)[:limit]
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Événements sécurité")
    header = [
        "Date (Africa/Algiers)",
        "Événement",
        "Risque",
        "Compte",
        "Identité saisie (empreinte)",
        "IP client",
        "IP proxy/technique",
        "Appareil",
        "Navigateur",
        "Système",
        "Route",
        "Session (empreinte)",
        "Motif",
        "Métadonnées non sensibles",
    ]
    header_cells = []
    for value in header:
        from openpyxl.cell import WriteOnlyCell

        cell = WriteOnlyCell(sheet, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
        header_cells.append(cell)
    sheet.append(header_cells)

    for event in events.iterator(chunk_size=1000):
        local_date = timezone.localtime(event.occurred_at).replace(tzinfo=None)
        sheet.append(
            [
                local_date,
                _excel_safe(event.get_event_type_display()),
                _excel_safe(event.get_risk_level_display()),
                _excel_safe(event.user.get_username() if event.user_id else ""),
                _excel_safe(event.submitted_identity_hash[:16] if event.submitted_identity_hash else ""),
                _excel_safe(event.client_ip),
                _excel_safe(event.proxy_ip),
                _excel_safe(event.device),
                _excel_safe(event.browser),
                _excel_safe(event.operating_system),
                _excel_safe(event.route),
                _excel_safe(event.session.session_fingerprint[:16] if event.session_id else ""),
                _excel_safe(event.reason),
                _excel_safe(json.dumps(event.metadata or {}, ensure_ascii=False, sort_keys=True)),
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    timestamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
    response["Content-Disposition"] = (
        f'attachment; filename="fbpack_evenements_securite_{timestamp}.xlsx"'
    )
    workbook.save(response)
    _record_admin_action(
        request,
        "SECURITY_EVENTS_EXPORTED",
        details=f"Export filtré, limite serveur {limit} lignes",
    )
    return response
