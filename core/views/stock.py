import openpyxl
import json
import traceback
import os
import re
import base64
import urllib.request
import urllib.error
from datetime import datetime, date
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db import models, transaction
from django.db.models import Sum, Q, F
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.models import User
from django.core.paginator import Paginator

from core.models import (
    Material, Supplier, SupplierContact, StockLocation, StockLot, StockMovement,
    DemandeAchat, BonCommande, LigneBonCommande, StockSeuil, Machine
)
try:
    from core.models import ConsommationEncre, ProductionOrder
except ImportError:
    pass
from core.forms import MaterialForm, SupplierForm
from .helpers import highlight_search

import logging
import math
from urllib.parse import quote

from core.decorators import module_required
from .security_utils import (
    parse_date, parse_float, parse_int, rate_limited, safe_json,
    self_approval_blocked, validate_upload, xl_safe,
)

logger = logging.getLogger('security')

LOT_STATUTS = ('EN_ATTENTE', 'CONFORME', 'BLOQUE', 'QUARANTAINE')
CERTIFICATE_EXTENSIONS = frozenset({'pdf', 'png', 'jpg', 'jpeg'})
CERTIFICATE_MAX_BYTES = 10 * 1024 * 1024
IMPORT_EXTENSIONS = frozenset({'xlsx', 'xlsm'})
IMPORT_MAX_BYTES = 10 * 1024 * 1024
IMPORT_MAX_ROWS = 20000
LABEL_IMAGE_TYPES = frozenset({'image/jpeg', 'image/png', 'image/webp'})
LABEL_MAX_BYTES = 6 * 1024 * 1024
_ISO_DATE = re.compile(r'^\d{4}-\d{2}-\d{2}$')

# ===========================================================================
# --- DROITS ET LOGIQUE DE SÉCURITÉ ---
# ===========================================================================

def can_manage_stock(user):
    return user.is_staff or user.groups.filter(name__in=['RESPONSABLE_STOCK', 'ADMIN']).exists()

def can_validate_purchase(user):
    return user.is_staff or user.groups.filter(name__in=['DIRECTION', 'ADMIN']).exists()


def _refuse_manage(request):
    messages.error(request, "Accès refusé.")
    return redirect('stock_advanced')


def _positive_float(value, label):
    number = parse_float(value)
    if number is None or number <= 0:
        raise ValueError(f"{label} invalide : une valeur positive est requise.")
    return number


def _non_negative_float(value, label):
    number = parse_float(value)
    if number is None or number < 0:
        raise ValueError(f"{label} invalide.")
    return number


def _clean_iso_date(value):
    """Date AAAA-MM-JJ valide, sinon '' (None et '' sont conservés tels quels)."""
    if not value:
        return value
    return value if _ISO_DATE.match(value) and parse_date(value) else ''


def _clean_int_str(value):
    if not value:
        return value
    return value if str(value).isdecimal() else ''


def _clean_ids(data):
    raw = data.getlist('ids') if hasattr(data, 'getlist') else data.get('ids', [])
    if not isinstance(raw, (list, tuple)):
        return []
    ids = []
    for value in raw[:1000]:
        number = parse_int(value, minimum=1)
        if number is not None:
            ids.append(number)
    return ids

# ===========================================================================
# --- VUE COMPATIBILITÉ HISTORIQUE STOCK_VIEW ---
# ===========================================================================

@login_required
def stock_view(request):
    """Redirection transparente ou affichage du module stock unifié."""
    return redirect('stock_advanced')

@login_required
@module_required('stock')
def clear_all_stock(request):
    if not request.user.is_superuser:
        messages.error(request, "Action réservée aux super-administrateurs.")
        return redirect('stock_advanced')
    if request.method == 'POST':
        try:
            with transaction.atomic():
                StockMovement.objects.all().delete()
                StockLot.objects.all().delete()
                count, _ = Material.objects.all().delete()
            logger.warning('stock_cleared user_id=%s deleted_rows=%s', request.user.id, count)
            messages.success(request, f"⚠️ Tout le stock a été vidé ({count} matières supprimées).")
        except Exception:
            logger.exception('stock_clear_failed')
            messages.error(request, "Erreur lors de la vidange. Détails dans les journaux du serveur.")
    return redirect('stock_advanced')

# ===========================================================================
# --- LOGIQUE TRANSACTIONNELLE SÉCURISÉE (SERVICES) ---
# ===========================================================================

class StockService:
    @staticmethod
    @transaction.atomic
    def enregistrer_mouvement(type_mvt, material, quantite, lot=None, src=None, dst=None, of=None, machine=None, user=None, motif=""):
        try:
            quantite_valide = math.isfinite(float(quantite)) and quantite > 0
        except (TypeError, ValueError):
            quantite_valide = False
        if not quantite_valide:
            raise ValueError("La quantité doit être strictement positive.")
        if type_mvt not in {code for code, _label in StockMovement.TYPE_CHOICES}:
            raise ValueError("Type de mouvement inconnu.")

        # Verrou des lignes concernées : deux mouvements simultanés ne s'écrasent plus.
        Material.objects.select_for_update().filter(pk=material.pk).first()
        material.refresh_from_db(fields=['quantity'])
        if lot:
            StockLot.objects.select_for_update().filter(pk=lot.pk).first()
            lot.refresh_from_db(fields=['quantite_restante'])

        if type_mvt in ['SORTIE', 'PERTE', 'RETOUR']:
            if lot and lot.quantite_restante < quantite:
                raise ValueError(f"Le lot {lot.numero_lot} n'a pas assez de stock disponible (Reste: {lot.quantite_restante} kg).")
            if material.usable_quantity < quantite:
                raise ValueError(f"La matière n'a pas assez de stock utilisable (Disponible: {material.usable_quantity} kg).")

        mvt = StockMovement.objects.create(
            type=type_mvt, material=material, lot=lot, quantite=quantite,
            emplacement_source=src, emplacement_destination=dst,
            of=of, machine=machine, utilisateur=user, motif=motif
        )

        if type_mvt == 'ENTREE':
            material.quantity = round(material.quantity + quantite, 3)
            if lot:
                lot.quantite_restante = round(lot.quantite_restante + quantite, 3)
        elif type_mvt in ['SORTIE', 'PERTE', 'RETOUR']:
            material.quantity = round(material.quantity - quantite, 3)
            if lot:
                lot.quantite_restante = round(lot.quantite_restante - quantite, 3)

        material.save()
        if lot:
            lot.save()
        return mvt

    @staticmethod
    @transaction.atomic
    def annuler_mouvement(mouvement, motif_annulation, user):
        # Verrou puis relecture : une annulation en double ne contre-passe plus deux fois.
        StockMovement.objects.select_for_update().filter(pk=mouvement.pk).first()
        Material.objects.select_for_update().filter(pk=mouvement.material_id).first()
        if mouvement.lot_id:
            StockLot.objects.select_for_update().filter(pk=mouvement.lot_id).first()
        mouvement.refresh_from_db()

        if mouvement.annule:
            raise ValueError("Ce mouvement est déjà annulé.")

        mouvement.annule = True
        mouvement.motif_annulation = motif_annulation
        mouvement.save()

        mat = mouvement.material
        lot = mouvement.lot
        qte = mouvement.quantite

        if mouvement.type == 'ENTREE':
            mat.quantity = round(mat.quantity - qte, 3)
            if lot:
                lot.quantite_restante = round(lot.quantite_restante - qte, 3)
        elif mouvement.type in ['SORTIE', 'PERTE', 'RETOUR']:
            mat.quantity = round(mat.quantity + qte, 3)
            if lot:
                lot.quantite_restante = round(lot.quantite_restante + qte, 3)

        mat.save()
        if lot:
            lot.save()


# ===========================================================================
# --- CONTROLE DE COHÉRENCE ---
# ===========================================================================

@login_required
@module_required('stock')
def controle_stock_view(request):
    materials = Material.objects.filter(is_archived=False)
    ecarts = []
    
    for m in materials:
        entrees = StockMovement.objects.filter(material=m, type='ENTREE', annule=False).aggregate(total=Sum('quantite'))['total'] or 0.0
        sorties = StockMovement.objects.filter(material=m, type__in=['SORTIE', 'PERTE', 'RETOUR'], annule=False).aggregate(total=Sum('quantite'))['total'] or 0.0
        
        theorique_ledger = m.initial_quantity + entrees - sorties
        somme_lots = m.lots.aggregate(total=Sum('quantite_restante'))['total'] or 0.0

        if abs(m.quantity - theorique_ledger) > 0.01 or abs(m.quantity - somme_lots) > 0.01:
            ecarts.append({
                'material': m,
                'stock_physique': m.quantity,
                'theorique_ledger': theorique_ledger,
                'somme_lots': somme_lots,
                'ecart': abs(m.quantity - somme_lots)
            })

    return render(request, 'stock/controle_stock.html', {'ecarts': ecarts})


@login_required
@module_required('stock')
def seuils_recalculer(request):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    materials = Material.objects.filter(is_archived=False)
    date_limite = timezone.now() - timezone.timedelta(days=30)
    recalculs = 0

    for m in materials:
        conso_totale = StockMovement.objects.filter(
            material=m, type='SORTIE', annule=False, date__gte=date_limite
        ).aggregate(total=Sum('quantite'))['total'] or 0.0
        
        conso_journaliere = round(conso_totale / 30.0, 2)
        
        seuil, _ = StockSeuil.objects.get_or_create(material=m)
        seuil.consommation_journaliere_moy = conso_journaliere
        seuil.save()
        recalculs += 1

    messages.success(request, f"🔮 Recalcul automatique terminé pour {recalculs} matière(s) sur les 30 derniers jours de production.")
    return redirect('stock_advanced')


# ===========================================================================
# --- VUES DES MATIÈRES PREMIÈRES (CRUD & ARCHIVE) ---
# ===========================================================================

@login_required
@module_required('stock')
def stock_advanced_view(request):
    try:
        tab = request.GET.get('tab', 'matieres')
        sub = request.GET.get('sub', 'da')
        search_query = request.GET.get('q', '').strip()
        category_filter = request.GET.get('category', '')
        low_stock_only = request.GET.get('low_stock', '') == 'on'
        supplier_filter = request.GET.get('supplier', '')
        sort_by = request.GET.get('sort', 'alert')
        archives = request.GET.get('archives', '') == '1'

        peut_gerer = can_manage_stock(request.user)
        peut_valider = can_validate_purchase(request.user)

        materials_list = Material.objects.select_related('supplier').filter(is_archived=archives)

        if search_query:
            materials_list = materials_list.filter(
                Q(name__icontains=search_query) |
                Q(code__icontains=search_query) |
                Q(supplier__name__icontains=search_query)
            )
        if category_filter:
            materials_list = materials_list.filter(category=category_filter)
        if supplier_filter:
            materials_list = materials_list.filter(supplier_id=supplier_filter)
        if low_stock_only:
            materials_list = [m for m in materials_list if m.is_low_stock()]

        if not isinstance(materials_list, list):
            if sort_by == 'name':
                materials_list = materials_list.order_by('name')
            elif sort_by == 'stock_asc':
                materials_list = materials_list.order_by('quantity')
            elif sort_by == 'stock_desc':
                materials_list = materials_list.order_by('-quantity')
            elif sort_by == 'price_desc':
                materials_list = materials_list.order_by('-price_per_unit')
            else:
                materials_list = sorted(materials_list, key=lambda x: (not x.is_low_stock(), x.usable_quantity))

        paginator_mat = Paginator(materials_list, 20)
        page_mat_num = request.GET.get('page_mat', 1)
        materials_page = paginator_mat.get_page(page_mat_num)

        all_active_materials = Material.objects.filter(is_archived=False)
        alertes_stock = []
        nb_ruptures = 0
        nb_critiques = 0
        nb_alertes_simples = 0

        cat_stats = {
            'FILM': {'rupture': 0, 'critique': 0, 'alerte': 0},
            'INK': {'rupture': 0, 'critique': 0, 'alerte': 0},
            'GLUE': {'rupture': 0, 'critique': 0, 'alerte': 0},
            'SOLV': {'rupture': 0, 'critique': 0, 'alerte': 0},
        }

        for m in all_active_materials:
            if m.is_low_stock():
                qty = m.usable_quantity
                thresh = m.min_threshold
                pct = round((qty / thresh) * 100, 1) if thresh > 0 else 0

                if qty <= 0:
                    niveau = 'RUPTURE'
                    icone = '🔴'
                    nb_ruptures += 1
                    if m.category in cat_stats:
                        cat_stats[m.category]['rupture'] += 1
                elif pct < 50:
                    niveau = 'CRITIQUE'
                    icone = '🟠'
                    nb_critiques += 1
                    if m.category in cat_stats:
                        cat_stats[m.category]['critique'] += 1
                else:
                    niveau = 'ALERTE'
                    icone = '🟡'
                    nb_alertes_simples += 1
                    if m.category in cat_stats:
                        cat_stats[m.category]['alerte'] += 1

                alertes_stock.append({
                    'id': m.id, 'name': m.name, 'category': m.category,
                    'cat_label': m.get_category_display(), 'quantity': qty, 'unit': m.unit,
                    'min_threshold': thresh, 'supplier': m.supplier.name if m.supplier else '—',
                    'pct': min(pct, 100), 'niveau': niveau, 'icone': icone,
                    'suggestion': max(0.0, (thresh * 2) - qty)
                })

        alertes_stock.sort(key=lambda x: (0 if x['niveau'] == 'RUPTURE' else (1 if x['niveau'] == 'CRITIQUE' else 2), -x['pct']))

        top_alertes = alertes_stock[:10]
        top_alertes_noms = [a['name'][:25] + '...' if len(a['name']) > 25 else a['name'] for a in top_alertes]
        top_alertes_stock = [a['quantity'] for a in top_alertes]
        top_alertes_seuil = [a['min_threshold'] for a in top_alertes]
        top_alertes_couleurs = []
        for a in top_alertes:
            if a['niveau'] == 'RUPTURE':
                top_alertes_couleurs.append('#dc2626')
            elif a['niveau'] == 'CRITIQUE':
                top_alertes_couleurs.append('#ea580c')
            else:
                top_alertes_couleurs.append('#ca8a04')

        cat_labels = ['Film/Papier', 'Encre', 'Colle', 'Solvant']
        cat_rupture = [cat_stats['FILM']['rupture'], cat_stats['INK']['rupture'], cat_stats['GLUE']['rupture'], cat_stats['SOLV']['rupture']]
        cat_critique = [cat_stats['FILM']['critique'], cat_stats['INK']['critique'], cat_stats['GLUE']['critique'], cat_stats['SOLV']['critique']]
        cat_alerte = [cat_stats['FILM']['alerte'], cat_stats['INK']['alerte'], cat_stats['GLUE']['alerte'], cat_stats['SOLV']['alerte']]

        lots_peremption = []
        lots_proches = StockLot.objects.filter(quantite_restante__gt=0, date_expiration__isnull=False).order_by('date_expiration')
        for lot in lots_proches:
            if lot.jours_avant_expiration <= 30:
                lots_peremption.append(lot)

        mouvements_list = StockMovement.objects.select_related('material', 'lot', 'emplacement_source', 'emplacement_destination', 'utilisateur').all()
        mvt_from = request.GET.get('mvt_from')
        mvt_to = request.GET.get('mvt_to')
        mvt_type = request.GET.get('mvt_type')
        mvt_cat = request.GET.get('mvt_cat')
        mvt_loc = request.GET.get('mvt_loc')
        mvt_from, mvt_to, mvt_loc = _clean_iso_date(mvt_from), _clean_iso_date(mvt_to), _clean_int_str(mvt_loc)
        mvt_q = request.GET.get('mvt_q')
        mvt_annules = request.GET.get('mvt_annules') == '1'

        if not mvt_annules:
            mouvements_list = mouvements_list.filter(annule=False)
        if mvt_from:
            mouvements_list = mouvements_list.filter(date__date__gte=mvt_from)
        if mvt_to:
            mouvements_list = mouvements_list.filter(date__date__lte=mvt_to)
        if mvt_type:
            mouvements_list = mouvements_list.filter(type=mvt_type)
        if mvt_cat:
            mouvements_list = mouvements_list.filter(material__category=mvt_cat)
        if mvt_loc:
            mouvements_list = mouvements_list.filter(Q(emplacement_source_id=mvt_loc) | Q(emplacement_destination_id=mvt_loc))
        if mvt_q:
            mouvements_list = mouvements_list.filter(Q(material__name__icontains=mvt_q) | Q(lot__numero_lot__icontains=mvt_q))

        paginator_mvt = Paginator(mouvements_list, 30)
        page_mvt_num = request.GET.get('page_mvt', 1)
        mouvements_page = paginator_mvt.get_page(page_mvt_num)

        previsions = []
        for m in all_active_materials:
            try:
                seuil = m.seuil_intelligent
                if seuil and seuil.consommation_journaliere_moy > 0:
                    jours = seuil.jours_de_stock
                    if jours <= 15:
                        previsions.append({
                            'material_id': m.id,
                            'material': m.name,
                            'stock_actuel': m.usable_quantity,
                            'conso_jour': seuil.consommation_journaliere_moy,
                            'jours_restants': jours,
                            'date_rupture': seuil.date_rupture_prevue,
                            'critique': jours <= 7
                        })
            except Exception:
                pass
        previsions.sort(key=lambda x: x['jours_restants'])

        suppliers_list = Supplier.objects.filter(is_archived=False).order_by('name')
        locations = StockLocation.objects.filter(is_active=True)
        lots_a_traiter = StockLot.objects.filter(statut__in=['EN_ATTENTE', 'BLOQUE', 'QUARANTAINE']).order_by('-date_reception')
        lots_tous = StockLot.objects.all().order_by('-id')[:100]
        
        da_statut = request.GET.get('da_statut', '')
        demandes = DemandeAchat.objects.all()
        if da_statut:
            demandes = demandes.filter(statut=da_statut)
            
        da_en_attente = DemandeAchat.objects.filter(statut='SOUMISE').count()
        bons_commande = BonCommande.objects.all().order_by('-date_commande')

        valeur_stock_total = sum(float(m.usable_quantity) * float(m.price_per_unit) for m in all_active_materials)

        context = {
            'tab': tab, 'sub': sub, 'search_query': search_query,
            'category_filter': category_filter, 'low_stock_only': low_stock_only,
            'supplier_filter': supplier_filter, 'sort': sort_by, 'archives': archives,
            
            'materials': materials_page, 
            'all_materials': all_active_materials,
            'materials_film': [m for m in all_active_materials if m.category == 'FILM'],
            'materials_ink': [m for m in all_active_materials if m.category == 'INK'],
            'materials_glue': [m for m in all_active_materials if m.category == 'GLUE'],
            'materials_solv': [m for m in all_active_materials if m.category == 'SOLV'],
            
            'total_matieres': all_active_materials.count(),
            'alertes_stock': alertes_stock, 'nb_alertes': len(alertes_stock),
            'nb_ruptures': nb_ruptures, 'nb_critiques': nb_critiques,
            'nb_alertes_simples': nb_alertes_simples, 'top_alertes': alertes_stock[:5],
            'top_alertes_noms': safe_json(top_alertes_noms),
            'top_alertes_stock': safe_json(top_alertes_stock),
            'top_alertes_seuil': safe_json(top_alertes_seuil),
            'top_alertes_couleurs': safe_json(top_alertes_couleurs),
            'cat_labels_json': safe_json(cat_labels),
            'cat_rupture_json': safe_json(cat_rupture),
            'cat_critique_json': safe_json(cat_critique),
            'cat_alerte_json': safe_json(cat_alerte),
            'previsions': previsions, 'lots_peremption': lots_peremption,
            'lots_a_traiter': lots_a_traiter, 'lots': lots_tous,
            'lots_bloques': StockLot.objects.filter(statut='BLOQUE').count(),
            'lots_attente': StockLot.objects.filter(statut='EN_ATTENTE').count(),
            'mouvements': mouvements_page, 'locations': locations, 'suppliers_list': suppliers_list,
            'suppliers': suppliers_list, 'demandes': demandes, 'da_en_attente': da_en_attente,
            'da_statut': da_statut, 'bons_commande': bons_commande,
            'valeur_stock_total': valeur_stock_total, 'categories': Material.CAT_CHOICES,
            'types_mouvement': StockMovement.TYPE_CHOICES,
            'peut_gerer': peut_gerer, 'peut_valider': peut_valider,
            'qs_mat': f"tab=matieres&q={quote(str(search_query))}&category={quote(str(category_filter))}&supplier={quote(str(supplier_filter))}&sort={quote(str(sort_by))}",
            'qs_mvt': f"tab=mouvements&mvt_from={quote(str(mvt_from or ''))}&mvt_to={quote(str(mvt_to or ''))}&mvt_type={quote(str(mvt_type or ''))}&mvt_cat={quote(str(mvt_cat or ''))}&mvt_loc={quote(str(mvt_loc or ''))}&mvt_q={quote(str(mvt_q or ''))}&mvt_annules={quote(request.GET.get('mvt_annules', ''))}"
        }
        return render(request, 'stock/stock_advanced.html', context)
    except Exception:
        logger.exception('stock_advanced_failed')
        messages.error(request, "❌ Erreur critique. Détails dans les journaux du serveur.")
        return redirect('dashboard')


@login_required
@module_required('stock')
def add_material(request):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    if request.method == 'POST':
        form = MaterialForm(request.POST)
        if form.is_valid():
            mat = form.save(commit=False)
            if not mat.initial_quantity and mat.quantity:
                mat.initial_quantity = mat.quantity
            mat.save()
            StockSeuil.objects.create(material=mat)
            messages.success(request, '✅ Matière ajoutée avec succès !')
            return redirect('stock_advanced')
    else:
        form = MaterialForm()
    return render(request, 'stock/material_form.html', {
        'form': form, 'title': 'Nouvelle Matière Première', 'btn_label': 'Ajouter la Matière'
    })


@login_required
@module_required('stock')
def edit_material(request, id):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    material = get_object_or_404(Material, id=id)
    if request.method == 'POST':
        form = MaterialForm(request.POST, instance=material)
        if form.is_valid():
            form.save()
            messages.success(request, f'✅ Matière "{material.name}" modifiée !')
            return redirect('stock_advanced')
    else:
        form = MaterialForm(instance=material)
    return render(request, 'stock/material_form.html', {
        'form': form, 'material': material, 'title': f'Modifier : {material.name}', 'btn_label': 'Enregistrer'
    })


@login_required
@module_required('stock')
def delete_material(request, id):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    material = get_object_or_404(Material, id=id)
    if request.method == 'POST':
        material.is_archived = True
        material.save()
        messages.success(request, f'📁 Matière « {material.name} » archivée avec succès.')
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def material_restore(request, id):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    material = get_object_or_404(Material, id=id)
    if request.method == 'POST':
        material.is_archived = False
        material.save()
        messages.success(request, f'♻️ Matière « {material.name} » restaurée parmi les matières actives.')
    return redirect('stock_advanced')


# ===========================================================================
# --- ACTIONS GROUPÉES SÉCURISÉES ---
# ===========================================================================

@login_required
@module_required('stock')
@transaction.atomic
def bulk_delete_materials(request):
    if not can_manage_stock(request.user):
        return JsonResponse({'status': 'error', 'message': 'Accès interdit.'}, status=403)

    if request.method == 'POST':
        try:
            data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
            ids = _clean_ids(data)
            if ids:
                materials = Material.objects.filter(id__in=ids)
                count = materials.count()
                materials.update(is_archived=True)
                return JsonResponse({'status': 'success', 'message': f'📁 {count} matière(s) archivée(s) avec succès.'})
            return JsonResponse({'status': 'error', 'message': 'Sélection vide.'})
        except ValueError as exc:
            transaction.set_rollback(True)
            return JsonResponse({'status': 'error', 'message': str(exc)})
        except Exception:
            transaction.set_rollback(True)
            logger.exception('stock_bulk_action_failed')
            return JsonResponse({'status': 'error', 'message': "Opération impossible. Détails dans les journaux du serveur."}, status=500)
    return redirect('stock_advanced')


@login_required
@module_required('stock')
@transaction.atomic
def bulk_delete_movements(request):
    if not can_manage_stock(request.user):
        return JsonResponse({'status': 'error', 'message': 'Accès interdit.'}, status=403)

    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            ids = _clean_ids(data)
            motif = str(data.get('motif', 'Annulation groupée')).strip()[:500]
            
            if ids:
                mouvements = StockMovement.objects.filter(id__in=ids, annule=False)
                count = mouvements.count()
                for m in mouvements:
                    StockService.annuler_mouvement(m, motif, request.user)
                return JsonResponse({'status': 'success', 'message': f'↩ {count} mouvement(s) annulé(s) (contre-passation effectuée).'})
            return JsonResponse({'status': 'error', 'message': 'Sélection vide.'})
        except ValueError as exc:
            transaction.set_rollback(True)
            return JsonResponse({'status': 'error', 'message': str(exc)})
        except Exception:
            transaction.set_rollback(True)
            logger.exception('stock_bulk_action_failed')
            return JsonResponse({'status': 'error', 'message': "Opération impossible. Détails dans les journaux du serveur."}, status=500)
    return redirect('stock_advanced')


# ===========================================================================
# --- ENREGISTREMENTS DE LOTS ET CONTROLE QUALITÉ (FEFO) ---
# ===========================================================================

@login_required
@module_required('stock')
@transaction.atomic
def lot_add(request):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    if request.method == 'POST':
        try:
            material_id = request.POST.get('material')
            material_name_free = request.POST.get('material_name_free', '').strip()
            category = request.POST.get('category')
            numero_lot = request.POST.get('numero_lot', '').strip()
            date_reception = request.POST.get('date_reception')
            fournisseur_id = request.POST.get('fournisseur') or None
            emplacement_id = request.POST.get('emplacement') or None
            quantite = _positive_float(request.POST.get('quantite_initiale'), 'Quantité')
            prix = _non_negative_float(request.POST.get('prix_unitaire', 0), 'Prix unitaire')
            statut = request.POST.get('statut', 'EN_ATTENTE')
            if statut not in LOT_STATUTS:
                raise ValueError("Statut de lot invalide.")
            notes = request.POST.get('notes', '')
            laize = request.POST.get('laize') or None
            longueur = request.POST.get('longueur') or None

            if not _clean_iso_date(date_reception):
                raise ValueError("Date de réception invalide.")
            certificat = request.FILES.get('certificat_qualite')
            if certificat:
                validate_upload(certificat, CERTIFICATE_EXTENSIONS, CERTIFICATE_MAX_BYTES)

            if material_id:
                material = get_object_or_404(Material, id=material_id)
            elif material_name_free:
                material, _ = Material.objects.get_or_create(
                    name=material_name_free,
                    defaults={'category': category or 'INK', 'quantity': 0, 'min_threshold': 100}
                )
            else:
                messages.error(request, "Veuillez désigner la matière première.")
                return redirect('stock_advanced')

            fournisseur = Supplier.objects.filter(id=fournisseur_id).first() if fournisseur_id else None
            emplacement = StockLocation.objects.filter(id=emplacement_id).first() if emplacement_id else None
            
            lot = StockLot.objects.create(
                material=material, numero_lot=numero_lot,
                date_reception=date_reception, fournisseur=fournisseur,
                emplacement=emplacement, quantite_initiale=quantite,
                quantite_restante=quantite, prix_unitaire=prix,
                statut=statut, notes=notes, laize=laize, longueur=longueur,
                created_by=request.user
            )

            if request.FILES.get('certificat_qualite'):
                lot.certificat_qualite = request.FILES['certificat_qualite']
                lot.save()

            if statut == 'CONFORME':
                StockService.enregistrer_mouvement(
                    type_mvt='ENTREE', material=material, quantite=quantite, lot=lot,
                    dst=emplacement, user=request.user, motif=f"Réception lot conforme {numero_lot}"
                )

            messages.success(request, f"✅ Lot {numero_lot} créé avec succès en Magasin Général.")
        except ValueError as exc:
            transaction.set_rollback(True)
            messages.error(request, f"❌ Erreur lors de la création du lot : {exc}")
        except Exception:
            transaction.set_rollback(True)
            logger.exception('stock_lot_add_failed')
            messages.error(request, "❌ Erreur lors de la création du lot. Détails dans les journaux du serveur.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
@transaction.atomic
def lot_valider(request, id):
    if not can_validate_purchase(request.user):
        messages.error(request, "Accès refusé (Droits de validation requis).")
        return redirect('stock_advanced')

    lot = get_object_or_404(StockLot, id=id)
    if request.method == 'POST':
        ancien_statut = lot.statut
        lot.statut = 'CONFORME'
        lot.save()

        deja_entre_en_stock = lot.mouvements.filter(type='ENTREE', annule=False).exists()
        if ancien_statut != 'CONFORME' and not deja_entre_en_stock and lot.quantite_restante > 0:
            StockService.enregistrer_mouvement(
                type_mvt='ENTREE', material=lot.material, quantite=lot.quantite_restante,
                lot=lot, dst=lot.emplacement, user=request.user, motif=f"Validation Qualité Lot {lot.numero_lot}"
            )
        messages.success(request, f"✅ Le lot {lot.numero_lot} est validé et utilisable en production.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
@transaction.atomic
def lot_bloquer(request, id):
    if not can_validate_purchase(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    lot = get_object_or_404(StockLot, id=id)
    if request.method == 'POST':
        lot.statut = 'BLOQUE'
        lot.save()
        messages.warning(request, f"🔴 Le lot {lot.numero_lot} a été bloqué et écarté du stock utilisable.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def lot_detail(request, id):
    lot = get_object_or_404(StockLot, id=id)
    mouvements = lot.mouvements.all().order_by('-date')
    return render(request, 'stock/lot_detail.html', {'lot': lot, 'mouvements': mouvements})


# ===========================================================================
# --- MOUVEMENTS UNITAIRES ---
# ===========================================================================

@login_required
@module_required('stock')
@transaction.atomic
def mouvement_add(request):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    if request.method == 'POST':
        try:
            material = get_object_or_404(Material, id=request.POST.get('material'))
            type_mvt = request.POST.get('type')
            quantite = _positive_float(request.POST.get('quantite'), 'Quantité')
            lot_id = request.POST.get('lot') or None
            src_id = request.POST.get('emplacement_source') or None
            dst_id = request.POST.get('emplacement_destination') or None
            of_id = request.POST.get('of') or None
            machine_id = request.POST.get('machine') or None
            motif = request.POST.get('motif', '')

            lot = StockLot.objects.filter(id=lot_id).first() if lot_id else None
            src = StockLocation.objects.filter(id=src_id).first() if src_id else None
            dst = StockLocation.objects.filter(id=dst_id).first() if dst_id else None
            machine_obj = Machine.objects.filter(id=machine_id).first() if machine_id else None

            if lot and lot.material_id != material.id:
                raise ValueError("Ce lot n'appartient pas à la matière sélectionnée.")

            if material.category == 'INK' and type_mvt == 'SORTIE':
                if not of_id or not machine_id:
                    raise ValueError("La traçabilité stricte exige qu'une sortie d'encre soit liée à une machine et un Ordre de Fabrication.")

            if not lot and type_mvt in ['SORTIE', 'PERTE']:
                lots_dispos = material.get_fefo_lots()
                if not lots_dispos.exists():
                    raise ValueError(f"Aucun lot conforme disponible pour « {material.name} ».")
                lot = lots_dispos.first()

            StockService.enregistrer_mouvement(
                type_mvt=type_mvt, material=material, quantite=quantite, lot=lot,
                src=src, dst=dst, machine=machine_obj, user=request.user, motif=motif
            )
            messages.success(request, f"🔄 Mouvement enregistré : {quantite} kg de {material.name}.")
        except ValueError as exc:
            transaction.set_rollback(True)
            messages.error(request, f"❌ Erreur : {exc}")
        except Exception:
            transaction.set_rollback(True)
            logger.exception('stock_movement_failed')
            messages.error(request, "❌ Erreur inattendue. Détails dans les journaux du serveur.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
@transaction.atomic
def mouvement_annuler(request, id):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    mvt = get_object_or_404(StockMovement, id=id)
    if request.method == 'POST':
        motif = request.POST.get('motif', '').strip()
        if not motif:
            messages.error(request, "Un motif est obligatoire pour annuler un mouvement.")
            return redirect('stock_advanced')
        try:
            StockService.annuler_mouvement(mvt, motif, request.user)
            messages.success(request, "✅ Mouvement annulé et contre-passé avec succès.")
        except ValueError as exc:
            transaction.set_rollback(True)
            messages.error(request, f"❌ Erreur : {exc}")
        except Exception:
            transaction.set_rollback(True)
            logger.exception('stock_movement_failed')
            messages.error(request, "❌ Erreur inattendue. Détails dans les journaux du serveur.")
    return redirect('stock_advanced')


# ===========================================================================
# --- CHAINE D'ACHATS SÉQUENTIELLE (DA -> BC -> RÉCEPTION) ---
# ===========================================================================

@login_required
@module_required('stock')
@transaction.atomic
def da_add(request):
    if request.method == 'POST':
        try:
            annee = timezone.now().year
            dernier = DemandeAchat.objects.filter(reference__startswith=f"DA-{annee}-").order_by('-reference').first()
            seq = 1
            if dernier:
                try:
                    seq = int(dernier.reference.split('-')[-1]) + 1
                except Exception:
                    pass
            ref = f"DA-{annee}-{seq:04d}"

            material = get_object_or_404(Material, id=request.POST.get('material'))
            DemandeAchat.objects.create(
                reference=ref, material=material,
                quantite_demandee=_positive_float(request.POST.get('quantite_demandee'), 'Quantité'),
                motif=request.POST.get('motif', ''),
                urgence=request.POST.get('urgence', 'NORMALE'),
                statut='SOUMISE', demandeur=request.user,
                date_besoin=_clean_iso_date(request.POST.get('date_besoin')) or None,
            )
            messages.success(request, f"📋 Demande d'achat {ref} soumise à la direction.")
        except ValueError as exc:
            transaction.set_rollback(True)
            messages.error(request, f"Erreur : {exc}")
        except Exception:
            transaction.set_rollback(True)
            logger.exception('stock_purchase_failed')
            messages.error(request, "Erreur inattendue. Détails dans les journaux du serveur.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def da_valider(request, id):
    if not can_validate_purchase(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    da = get_object_or_404(DemandeAchat, id=id)
    if request.method == 'POST':
        if da.statut != 'SOUMISE':
            messages.warning(request, f"DA {da.reference} : déjà traitée.")
            return redirect('stock_advanced')
        if self_approval_blocked(request.user, da.demandeur_id):
            messages.error(request, "Vous ne pouvez pas valider votre propre demande d'achat.")
            return redirect('stock_advanced')
        da.statut = 'VALIDEE'
        da.valideur = request.user
        da.date_validation = timezone.now().date()
        da.save()
        messages.success(request, f"DA {da.reference} validée par la direction.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def da_refuser(request, id):
    if not can_validate_purchase(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    da = get_object_or_404(DemandeAchat, id=id)
    if request.method == 'POST':
        if da.statut not in ('SOUMISE', 'VALIDEE'):
            messages.warning(request, f"DA {da.reference} : déjà traitée.")
            return redirect('stock_advanced')
        da.statut = 'REFUSEE'
        da.save()
        messages.warning(request, f"DA {da.reference} refusée.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
@transaction.atomic
def bc_add(request):
    if not can_manage_stock(request.user):
        messages.error(request, "Accès refusé.")
        return redirect('stock_advanced')

    if request.method == 'POST':
        try:
            annee = timezone.now().year
            dernier = BonCommande.objects.filter(reference__startswith=f"BC-{annee}-").order_by('-reference').first()
            seq = 1
            if dernier:
                try:
                    seq = int(dernier.reference.split('-')[-1]) + 1
                except Exception:
                    pass
            ref = f"BC-{annee}-{seq:04d}"

            fournisseur = get_object_or_404(Supplier, id=request.POST.get('fournisseur'))
            bc = BonCommande.objects.create(
                reference=ref, fournisseur=fournisseur, statut='BROUILLON',
                date_livraison_prevue=_clean_iso_date(request.POST.get('date_livraison_prevue')) or None,
                notes=request.POST.get('notes', ''), cree_par=request.user,
            )
            idx = 1
            total = 0
            while idx <= 200 and request.POST.get(f'material_{idx}'):
                mat = Material.objects.filter(id=request.POST.get(f'material_{idx}')).first()
                if mat:
                    qte = _positive_float(request.POST.get(f'quantite_{idx}'), 'Quantité')
                    prix = _non_negative_float(request.POST.get(f'prix_{idx}', 0), 'Prix')
                    LigneBonCommande.objects.create(
                        bon_commande=bc, material=mat,
                        quantite_commandee=qte, prix_unitaire=prix
                    )
                    total += qte * prix
                idx += 1
            
            da_id = parse_int(request.GET.get('da'))
            if da_id:
                da = DemandeAchat.objects.filter(id=da_id).first()
                # Seule une demande déjà validée par la direction peut devenir un bon de commande.
                if da and da.statut == 'VALIDEE':
                    da.statut = 'COMMANDEE'
                    da.bon_commande = bc
                    da.save()

            bc.montant_total = total
            bc.save()
            messages.success(request, f"📄 BC {bc.reference} généré avec succès.")
        except ValueError as exc:
            transaction.set_rollback(True)
            messages.error(request, f"Erreur : {exc}")
        except Exception:
            transaction.set_rollback(True)
            logger.exception('stock_purchase_failed')
            messages.error(request, "Erreur inattendue. Détails dans les journaux du serveur.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def bc_envoyer(request, id):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    bc = get_object_or_404(BonCommande, id=id)
    if request.method == 'POST':
        if bc.statut != 'BROUILLON':
            messages.warning(request, f"BC {bc.reference} : déjà envoyé ou clôturé.")
            return redirect('stock_advanced')
        bc.statut = 'ENVOYE'
        bc.save()
        messages.success(request, f"BC {bc.reference} marqué comme envoyé au fournisseur.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
@transaction.atomic
def bc_reception(request, id):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    bc = get_object_or_404(BonCommande.objects.select_for_update(), id=id)
    if request.method == 'POST':
        # Une réception répétée (double clic, rechargement) ne recrée plus les lots.
        if bc.statut == 'RECU_TOTAL':
            messages.warning(request, f"BC {bc.reference} : déjà réceptionné.")
            return redirect('stock_advanced')
        bc.statut = 'RECU_TOTAL'
        bc.date_livraison_reelle = timezone.now().date()
        bc.save()
        
        for ligne in bc.lignes.all():
            num_lot = f"LOT-{bc.reference}-{ligne.material.id}"
            StockLot.objects.create(
                material=ligne.material, numero_lot=num_lot,
                date_reception=timezone.now().date(),
                fournisseur=bc.fournisseur,
                quantite_initiale=ligne.quantite_commandee,
                quantite_restante=ligne.quantite_commandee,
                prix_unitaire=ligne.prix_unitaire,
                statut='EN_ATTENTE', created_by=request.user,
            )
            ligne.quantite_recue = ligne.quantite_commandee
            ligne.save()
        messages.success(request, f"📦 BC {bc.reference} réceptionné : lots placés en attente de validation qualité.")
    return redirect('stock_advanced')


# ===========================================================================
# --- AUTRES CONFIGURATIONS, EXPORTS ET GRAPHES ---
# ===========================================================================

@login_required
@module_required('stock')
def seuil_update(request, material_id):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    material = get_object_or_404(Material, id=material_id)
    if request.method == 'POST':
        seuil, _ = StockSeuil.objects.get_or_create(
            material=material,
            defaults={'consommation_journaliere_moy': 0, 'delai_fournisseur_jours': 7}
        )
        conso = parse_float(request.POST.get('conso_jour', 0), minimum=0)
        delai = parse_int(request.POST.get('delai_jours', 7), minimum=0, maximum=3650)
        securite = parse_int(request.POST.get('securite_jours', 3), minimum=0, maximum=3650)
        if conso is None or delai is None or securite is None:
            messages.error(request, "❌ Valeurs invalides.")
            return redirect('stock_advanced')
        seuil.consommation_journaliere_moy = conso
        seuil.delai_fournisseur_jours = delai
        seuil.stock_securite_jours = securite
        seuil.save()
        messages.success(request, f"🔮 Paramètres prévisionnels mis à jour pour {material.name}.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def location_add(request):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        type_loc = request.POST.get('type', 'GENERAL')
        desc = request.POST.get('description', '')
        if name:
            StockLocation.objects.create(name=name, type=type_loc, description=desc)
            messages.success(request, f"Emplacement « {name} » créé.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def location_delete(request, id):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    loc = get_object_or_404(StockLocation, id=id)
    if request.method == 'POST':
        loc.delete()
        messages.success(request, "Emplacement supprimé.")
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def add_supplier(request):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        email = request.POST.get('email', '').strip()
        if name:
            supplier = Supplier.objects.create(name=name, email=email)
            contact_noms = request.POST.getlist('contact_nom[]')
            contact_postes = request.POST.getlist('contact_poste[]')
            contact_telephones = request.POST.getlist('contact_telephone[]')
            contact_emails = request.POST.getlist('contact_email[]')
            for i in range(len(contact_noms)):
                nom = contact_noms[i].strip() if i < len(contact_noms) else ''
                if nom:
                    SupplierContact.objects.create(
                        supplier=supplier,
                        nom=nom,
                        poste=contact_postes[i].strip() if i < len(contact_postes) else '',
                        telephone=contact_telephones[i].strip() if i < len(contact_telephones) else '',
                        email=contact_emails[i].strip() if i < len(contact_emails) else ''
                    )
            messages.success(request, '✅ Fournisseur ajouté avec succès !')
            return redirect('stock_advanced')
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def edit_supplier(request, id):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    supplier = get_object_or_404(Supplier, id=id)
    if request.method == 'POST':
        supplier.name = request.POST.get('name', supplier.name)
        supplier.email = request.POST.get('email', supplier.email)
        supplier.save()
        messages.success(request, f'✅ Fournisseur "{supplier.name}" modifié !')
        return redirect('stock_advanced')
    return redirect('stock_advanced')


@login_required
@module_required('stock')
def delete_supplier(request, id):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    supplier = get_object_or_404(Supplier, id=id)
    if request.method == 'POST':
        supplier.is_archived = True
        supplier.save()
        messages.success(request, f'🗑️ Fournisseur "{supplier.name}" archivé.')
    return redirect('stock_advanced')


@login_required
@module_required('stock', 'dashboard')
def stock_dashboard_data(request):
    critiques = []
    for m in Material.objects.filter(is_archived=False):
        if m.is_low_stock():
            critiques.append({
                'name': m.name, 
                'stock': m.usable_quantity, 
                'seuil': m.min_threshold
            })
    return JsonResponse({'critiques': critiques})


@login_required
@module_required('stock', 'crm', 'prepress', 'production', 'planification')
def material_search_api(request):
    query = request.GET.get('q', '').strip()
    if len(query) < 1:
        return JsonResponse([], safe=False)

    materials = Material.objects.filter(is_archived=False).filter(
        Q(name__icontains=query) | Q(code__icontains=query)
    )[:10]

    results = []
    for m in materials:
        results.append({
            'id': m.id,
            'name': m.name,
            'code': m.code,
            'stock': m.usable_quantity,
            'unit': m.unit,
            'thickness': m.thickness_or_grammage_display,
            'metrage': m.metrage_display
        })
    return JsonResponse(results, safe=False)


# ===========================================================================
# --- STATION D'IMPORTATION EXCEL ET FORMATS ---
# ===========================================================================

@login_required
@module_required('import', 'stock')
def import_stock_view(request):
    if not can_manage_stock(request.user):
        return _refuse_manage(request)
    crm_users = User.objects.filter(is_active=True).order_by('first_name', 'username')
    if request.method != 'POST':
        return render(request, 'stock/import_stock.html', {'crm_users': crm_users})

    import_type = request.POST.get('import_type', 'MOUVEMENTS')
    excel_file = request.FILES.get('excel_file')

    if not excel_file:
        messages.error(request, "❌ Veuillez sélectionner un fichier Excel.")
        return render(request, 'stock/import_stock.html', {'crm_users': crm_users})

    try:
        validate_upload(excel_file, IMPORT_EXTENSIONS, IMPORT_MAX_BYTES)
    except ValueError as exc:
        messages.error(request, f"❌ {exc}")
        return render(request, 'stock/import_stock.html', {'crm_users': crm_users})

    details = []
    success_count = 0
    error_count = 0

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        ws = wb.worksheets[0]
        details.append(f"📄 Analyse : « {ws.title} » ({ws.max_row} lignes)")
        if ws.max_row > IMPORT_MAX_ROWS:
            raise ValueError(f"Fichier trop volumineux : {IMPORT_MAX_ROWS} lignes maximum.")

        if import_type == 'MOUVEMENTS':
            for row_idx in range(2, ws.max_row + 1):
                raw_date = ws.cell(row=row_idx, column=1).value
                raw_machine = ws.cell(row=row_idx, column=2).value
                raw_produit = ws.cell(row=row_idx, column=3).value
                raw_quantite = ws.cell(row=row_idx, column=4).value
                raw_code = ws.cell(row=row_idx, column=5).value
                raw_notes = ws.cell(row=row_idx, column=6).value

                if not raw_produit and not raw_code:
                    continue

                try:
                    qte_val = abs(float(raw_quantite or 0))
                    if qte_val == 0:
                        continue
                    
                    mat = Material.objects.filter(code=raw_code).first() or Material.objects.filter(name=raw_produit).first()
                    if not mat:
                        mat = Material.objects.create(
                            name=raw_produit or f"Import-{raw_code}", code=raw_code or '',
                            category='INK', quantity=0, min_threshold=50
                        )
                        details.append(f"Ligne {row_idx} : 🆕 Nouvelle matière créée « {mat.name} »")

                    StockService.enregistrer_mouvement(
                        type_mvt='SORTIE', material=mat, quantite=qte_val,
                        user=request.user, motif=f"Import Excel : {raw_notes or ''}"
                    )
                    success_count += 1
                except Exception as row_err:
                    error_count += 1
                    details.append(f"Ligne {row_idx} : ❌ Erreur ({str(row_err)})")

        elif import_type == 'STOCK':
            for row_idx in range(2, ws.max_row + 1):
                designation = ws.cell(row=row_idx, column=1).value
                code_val = ws.cell(row=row_idx, column=2).value
                fournisseur_name = ws.cell(row=row_idx, column=3).value
                category = str(ws.cell(row=row_idx, column=4).value or '')
                
                # --- LECTURE DES COLONNES "Micronage / Grammage" (col 5) ET "Metrage (m)" (col 6) ---
                mic_grm_val = ws.cell(row=row_idx, column=5).value
                metrage = ws.cell(row=row_idx, column=6).value
                
                stock_initial = float(ws.cell(row=row_idx, column=7).value or 0)
                qty_val = float(ws.cell(row=row_idx, column=8).value or stock_initial)
                unit = ws.cell(row=row_idx, column=9).value or 'kg'
                seuil = float(ws.cell(row=row_idx, column=10).value or 50)
                prix = float(ws.cell(row=row_idx, column=11).value or 0)

                if not designation:
                    continue

                try:
                    supplier_obj = None
                    if fournisseur_name:
                        supplier_obj, _ = Supplier.objects.get_or_create(name=fournisseur_name)

                    mic_val = None
                    gram_val = None
                    if mic_grm_val is not None:
                        desig_upper = str(designation).upper()
                        cat_upper = str(category).upper()
                        if 'PAPIER' in desig_upper or 'KRAFT' in desig_upper or 'COUCH' in desig_upper or 'PAPIER' in cat_upper:
                            gram_val = float(mic_grm_val)
                        else:
                            mic_val = int(float(mic_grm_val))

                    met_val = float(metrage) if metrage else None
                    cat_code = 'FILM' if 'FILM' in str(category).upper() or 'PAPIER' in str(category).upper() else 'INK'

                    mat, created = Material.objects.get_or_create(
                        code=code_val,
                        defaults={
                            'name': designation, 
                            'category': cat_code,
                            'micronage_standard': mic_val,
                            'grammage': gram_val,
                            'metrage_standard': met_val,
                            'quantity': qty_val, 
                            'initial_quantity': stock_initial, 
                            'unit': unit,
                            'min_threshold': seuil, 
                            'price_per_unit': prix, 
                            'supplier': supplier_obj
                        }
                    )
                    if not created:
                        mat.quantity = qty_val
                        if mic_val is not None: mat.micronage_standard = mic_val
                        if gram_val is not None: mat.grammage = gram_val
                        if met_val is not None: mat.metrage_standard = met_val
                        mat.save()

                    success_count += 1
                except Exception as e:
                    error_count += 1
                    details.append(f"Ligne {row_idx} : ❌ {str(e)}")

        messages.success(request, f"🎉 Importation réussie : {success_count} ligne(s) ajoutée(s). {error_count} erreur(s).")
        return render(request, 'stock/import_stock.html', {'success': True, 'details': details, 'crm_users': crm_users})
    except ValueError as exc:
        return render(request, 'stock/import_stock.html', {'success': False, 'message': str(exc), 'crm_users': crm_users})
    except Exception:
        logger.exception('stock_import_failed')
        return render(request, 'stock/import_stock.html', {
            'success': False,
            'message': "Fichier illisible ou invalide. Détails dans les journaux du serveur.",
            'crm_users': crm_users,
        })


@login_required
@module_required('import', 'stock')
def download_template_stock(request):
    """Génère le bon template Excel à jour pour le Stock Matières"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template Stock"
    
    # En-têtes exacts
    ws.append([
        'Designation', 'Code', 'Fournisseur', 'Categorie (Film/Encre/Colle/Solvant)', 
        'Micronage / Grammage', 'Metrage (m)', 'Stock_Initial', 'Stock_Reel', 
        'Unite', 'Seuil_Min', 'Prix_Unitaire'
    ])
    
    # Exemple 1 : Film avec Micronage
    ws.append(['BOPP TRANSPARENT 20UM', 'BOPP20', 'SunChemical', 'Film/Papier', 20, 6000, 500, 500, 'kg', 100, 1200])
    # Exemple 2 : Papier avec Grammage
    ws.append(['PAPIER KRAFT BLANCHI 70G', 'KRAFT70', 'JPR', 'Film/Papier', 70, 5000, 1000, 1000, 'kg', 200, 1500])
    # Exemple 3 : Encre (sans micronage ni métrage)
    ws.append(['ENCRE BLUE CYAN HP RG', '03.043.CX.SVR', 'Chemigold', 'Encre', '', '', 50, 50, 'kg', 20, 2500])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="template_stock_matieres.xlsx"'
    wb.save(response)
    return response


@login_required
@module_required('import', 'stock', 'production', 'prepress')
def download_template_special_prod(request):
    """Génère le bon template Excel complet à jour pour la Production Spéciale"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template Prod"
    
    # En-têtes complets (18 colonnes)
    ws.append([
        'Date', 'Produit', 'Support', 'Qte_Lancee', 'Lot', 'Laize', 'Client', 
        'Equipe (A/B/C)', 'Machine', 'H_Debut', 'H_Fin', 'Prod_ML', 
        'Dec_Demarrage', 'Dec_Lisiere', 'Dec_Jonction', 'Dec_Transport', 
        'Prod_KG', 'Rebobinage_KG'
    ])
    
    # Ligne d'exemple
    ws.append([
        '2025-01-15', 'Sachet Pâtes 1Kg', 'BOPP 20', 500, 'L2025-01', 820, 
        'Client X', 'A', 'Flexo 1', '08:00', '16:00', 4500, 
        2.5, 3.0, 0.5, 1.0, 480, 0
    ])
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="template_production.xlsx"'
    wb.save(response)
    return response


@login_required
@module_required('stock')
def export_search_results(request):
    materials = Material.objects.filter(is_archived=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(['Désignation', 'Code', 'Catégorie', 'Micronage / Grammage', 'Métrage (m)', 'Stock Réel', 'Unité', 'Seuil Min', 'Prix/U', 'Valeur'])
    for m in materials:
        ws.append([
            xl_safe(m.name), xl_safe(m.code), m.get_category_display(),
            m.thickness_or_grammage_display, m.metrage_display,
            m.usable_quantity, xl_safe(m.unit), m.min_threshold, m.price_per_unit,
            m.usable_quantity * float(m.price_per_unit)
        ])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="stock_matiere.xlsx"'
    wb.save(response)
    return response


@login_required
@module_required('stock')
def export_mouvements(request):
    mouvements = StockMovement.objects.select_related('material', 'lot', 'emplacement_source', 'emplacement_destination', 'utilisateur').all()
    mvt_from = request.GET.get('mvt_from')
    mvt_to = request.GET.get('mvt_to')
    mvt_type = request.GET.get('mvt_type')
    mvt_cat = request.GET.get('mvt_cat')
    mvt_loc = request.GET.get('mvt_loc')
    mvt_from, mvt_to, mvt_loc = _clean_iso_date(mvt_from), _clean_iso_date(mvt_to), _clean_int_str(mvt_loc)
    mvt_q = request.GET.get('mvt_q')
    mvt_annules = request.GET.get('mvt_annules') == '1'

    if not mvt_annules:
        mouvements = mouvements.filter(annule=False)
    if mvt_from:
        mouvements = mouvements.filter(date__date__gte=mvt_from)
    if mvt_to:
        mouvements = mouvements.filter(date__date__lte=mvt_to)
    if mvt_type:
        mouvements = mouvements.filter(type=mvt_type)
    if mvt_cat:
        mouvements = mouvements.filter(material__category=mvt_cat)
    if mvt_loc:
        mouvements = mouvements.filter(Q(emplacement_source_id=mvt_loc) | Q(emplacement_destination_id=mvt_loc))
    if mvt_q:
        mouvements = mouvements.filter(Q(material__name__icontains=mvt_q) | Q(lot__numero_lot__icontains=mvt_q))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mouvements Stock"
    ws.append(['Date', 'Type', 'Matière', 'Code', 'Lot', 'Quantité (kg)', 'Source', 'Destination', 'Machine', 'OF', 'Utilisateur', 'Motif', 'Annulé'])
    for mv in mouvements:
        ws.append([
            mv.date.strftime('%d/%m/%Y %H:%M') if mv.date else '',
            mv.get_type_display(),
            xl_safe(mv.material.name) if mv.material else '',
            xl_safe(mv.material.code) if mv.material else '',
            xl_safe(mv.lot.numero_lot) if mv.lot else '',
            mv.quantite,
            xl_safe(mv.emplacement_source.name) if mv.emplacement_source else '',
            xl_safe(mv.emplacement_destination.name) if mv.emplacement_destination else '',
            xl_safe(mv.machine.name) if mv.machine else '',
            xl_safe(str(mv.of)) if mv.of else '',
            xl_safe(mv.utilisateur.username) if mv.utilisateur else '',
            xl_safe(mv.motif or ''),
            'OUI' if mv.annule else 'NON'
        ])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="journal_mouvements.xlsx"'
    wb.save(response)
    return response


# ===========================================================================
# --- API SCANNER IA (AVEC SÉCURISATION & JOURNALISATION) ---
# ===========================================================================

@login_required
@module_required('stock')
def scan_label_ai(request):
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'Requête POST requise.'}, status=400)

    if rate_limited(f'scan-label:{request.user.id}', 10, 60):
        return JsonResponse({'status': 'error', 'message': 'Trop de lectures, patientez une minute.'}, status=429)

    image_file = request.FILES.get('image')
    if not image_file:
        return JsonResponse({'status': 'error', 'message': 'Image manquante.'}, status=400)
    if image_file.size > LABEL_MAX_BYTES:
        return JsonResponse({'status': 'error', 'message': 'Image trop volumineuse (6 Mo maximum).'}, status=413)
    if (image_file.content_type or '') not in LABEL_IMAGE_TYPES:
        return JsonResponse({'status': 'error', 'message': 'Format refusé (JPEG, PNG ou WEBP).'}, status=400)

    api_key = os.environ.get('GEMINI_API_KEY', '').strip()
    if not api_key:
        return JsonResponse({'status': 'error', 'message': 'Module IA non configuré (Clé Gemini manquante).'}, status=400)

    try:
        image_bytes = image_file.read()
        image_b64 = base64.b64encode(image_bytes).decode('utf-8')
        mime_type = image_file.content_type or 'image/jpeg'

        prompt = """
Analyse l'étiquette industrielle et renvoie UNIQUEMENT un JSON structuré :
{
  "fournisseur": "SunChemical" ou "GulfPack" ou "JPR" ou null,
  "material_name": "string complet",
  "category": "FILM" ou "INK" ou "GLUE" ou "SOLV",
  "numero_lot": "string ou null",
  "poids_net": number_kg ou null,
  "laize_width": number_mm ou null,
  "longueur_length": number_m ou null
}
Renvoie uniquement le format brut JSON, pas de balises markdown, pas de texte d'explication.
"""
        payload = {
            "contents": [{
                "parts": [
                    {"text": prompt},
                    {"inline_data": {"mime_type": mime_type, "data": image_b64}}
                ]
            }]
        }

        url = "https://generativelanguage.googleapis.com/v1beta/models/gemini-3.8-flash:generateContent"
        req = urllib.request.Request(
            url, data=json.dumps(payload).encode('utf-8'),
            headers={'Content-Type': 'application/json', 'x-goog-api-key': api_key}
        )
        
        with urllib.request.urlopen(req, timeout=30) as resp:
            res_body = json.loads(resp.read().decode('utf-8'))
            text_response = res_body['candidates'][0]['content']['parts'][0]['text'].strip()
            text_clean = re.sub(r'```json\s*|\s*```', '', text_response).strip().strip('`').strip()
            parsed_json = json.loads(text_clean)

            return JsonResponse({'status': 'success', 'data': parsed_json})
    except Exception:
        logger.exception('stock_scan_label_failed')
        return JsonResponse({'status': 'error', 'message': "Erreur IA. Détails dans les journaux du serveur."}, status=500)


# ===========================================================================
# --- STUBS COMPATIBILITÉ HISTORIQUE ---
# ===========================================================================

@login_required
def add_consommation(request):
    return redirect('stock_advanced')

@login_required
def conso_list_view(request):
    return redirect('stock_advanced')

@login_required
def location_list(request):
    return redirect('stock_advanced')

@login_required
def lot_list(request):
    return redirect('stock_advanced')