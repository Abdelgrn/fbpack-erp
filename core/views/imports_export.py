import os
import re
import datetime
import unicodedata
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import timedelta

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.core.files.storage import FileSystemStorage
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.db.models import Q
from django.contrib.auth.models import User

from ..models import (
    Client, TechnicalProduct, Tooling, Quote, ProductionOrder, Machine,
    Material, ConsommationEncre, ProductionEntry
)


# ===========================================================================
# --- HELPER : Queryset des utilisateurs ayant accès au module CRM ---
# ===========================================================================

def get_crm_users_queryset():
    """
    Retourne uniquement les utilisateurs actifs ayant la permission 
    d'accéder au module CRM (ou les super-administrateurs).
    """
    return User.objects.filter(
        Q(is_active=True) & 
        (Q(is_superuser=True) | Q(module_permissions__can_access_crm=True))
    ).distinct().order_by('first_name', 'username')


def parse_time_safe(val):
    if val is None or val == 0 or val == '' or str(val).strip() == '':
        return None
    try:
        if hasattr(val, 'hour'):
            return val
        if hasattr(val, 'time'):
            return val.time()
        if isinstance(val, float):
            total_seconds = int(val * 24 * 3600)
            return datetime.time(total_seconds // 3600, (total_seconds % 3600) // 60)
        s = str(val).strip()
        parts = s.replace('h', ':').replace('H', ':').split(':')
        if len(parts) >= 2:
            return datetime.time(int(parts[0]), int(parts[1]))
        return None
    except Exception:
        return None

def normalize_header(h):
    if h is None:
        return ''
    h = str(h).strip()
    h = unicodedata.normalize('NFKD', h).encode('ascii', 'ignore').decode('ascii')
    h = h.lower()
    h = re.sub(r'[\s_]+', '_', h)
    h = h.strip('_')
    return h

COLUMN_CANDIDATES = {
    'Designation': ['designation', 'nom', 'name', 'produit', 'article', 'libelle'],
    'Categorie': ['categorie', 'category', 'type', 'famille', 'type_produit'],
    'Quantite': ['quantite', 'qte', 'stock', 'quantity', 'qte_stock'],
    'Unite': ['unite', 'unit', 'uom'],
    'Seuil_Min': ['seuil_min', 'seuil', 'min', 'seuil_minimum', 'stock_min', 'stock_minimum'],
    'Prix': ['prix', 'prix_unitaire', 'price', 'prix_unite', 'pu'],
}

def build_column_rename_map(columns):
    rename_map = {}
    used_canonicals = set()
    for col in columns:
        norm = normalize_header(col)
        for canonical, aliases in COLUMN_CANDIDATES.items():
            if canonical in used_canonicals:
                continue
            if norm in aliases:
                rename_map[col] = canonical
                used_canonicals.add(canonical)
                break
    return rename_map

CATEGORY_KEYWORDS = {
    'INK': ['encre', 'ink', 'tinta'],
    'SOLV': ['solvant', 'solv', 'diluant', 'thinner'],
    'GLUE': ['colle', 'glue', 'adhesif', 'adhesive'],
    'FILM': ['film', 'papier', 'support', 'paper', 'bobine'],
}

def _text_matches_category(text):
    for code, keywords in CATEGORY_KEYWORDS.items():
        if any(k in text for k in keywords):
            return code
    return None

def detect_material_category(raw_cat, designation):
    def _clean(t):
        t = unicodedata.normalize('NFKD', str(t or '')).encode('ascii', 'ignore').decode('ascii')
        return t.lower()

    raw_cat_clean = _clean(raw_cat)
    designation_clean = _clean(designation)

    code = _text_matches_category(raw_cat_clean)
    if code:
        return code, 'colonne Catégorie'

    code = _text_matches_category(designation_clean)
    if code:
        return code, 'nom du produit (fallback)'

    return 'FILM', 'valeur par défaut (aucune correspondance trouvée)'


# ===========================================================================
# --- HELPERS IMPORT CRM (MAPPING STRICT FICHIER EXCEL → ERP) ---
# ===========================================================================

def safe_str(val):
    if val is None:
        return ''
    try:
        if pd.isna(val):
            return ''
    except Exception:
        pass
    s = str(val).strip()
    if s.lower() in ('0', '0.0', 'nan', 'none', 'nat', ''):
        return ''
    if isinstance(val, float) and val == int(val):
        s = str(int(val))
    return s


def auto_detect_crm_header(df):
    target_keywords = ['nom_client', 'nom client', 'id_client', 'id client', 'raison_sociale']

    current_cols = [normalize_header(c) for c in df.columns]
    if any(k.replace(' ', '_') in current_cols or k in current_cols for k in target_keywords):
        return df

    for idx, row in df.head(15).iterrows():
        row_values = [normalize_header(v) for v in row.values]
        joined = ' | '.join(row_values)
        if ('nom_client' in row_values or 'nom client' in joined) and (
            'id_client' in row_values or 'id client' in joined or 'ville' in row_values
        ):
            new_headers = [str(v).strip() if str(v).strip() not in ('', 'nan', 'None') else f'COL_{i}'
                           for i, v in enumerate(row.values)]
            df_reheaded = df.iloc[idx + 1:].copy()
            df_reheaded.columns = new_headers
            df_reheaded = df_reheaded.reset_index(drop=True)
            return df_reheaded

    return df


def build_crm_column_index(columns):
    index = {}
    for col in columns:
        norm = normalize_header(col)
        if norm and norm not in index:
            index[norm] = col
    return index


def crm_get(col_index, row, aliases, default=''):
    for alias in aliases:
        a = normalize_header(alias)
        if a in col_index:
            val = safe_str(row.get(col_index[a], ''))
            if val:
                return val
    for alias in aliases:
        a = normalize_header(alias)
        if len(a) < 4:
            continue
        for norm, orig in col_index.items():
            if norm == a or norm.startswith(a + '_') or norm.endswith('_' + a):
                val = safe_str(row.get(orig, ''))
                if val:
                    return val
    return default


def map_statut_compte(raw):
    s = safe_str(raw).upper()
    if not s:
        return 'PROSPECT'
    if 'VIP' in s:
        return 'VIP'
    if 'INACTIF' in s or 'INACTIVE' in s or 'PERDU' in s or 'LOST' in s:
        return 'LOST'
    if 'ACTIF' in s or 'ACTIVE' in s:
        return 'ACTIVE'
    if 'PROSPECT' in s:
        return 'PROSPECT'
    return 'PROSPECT'


def map_conditions_paiement(raw):
    s = safe_str(raw).upper().replace(' ', '')
    if not s:
        return ''
    mapping = {
        'COMPTANT': 'COMPTANT', '0': 'COMPTANT',
        '30': '30J', '30J': '30J', '30JOURS': '30J',
        '45': '45J', '45J': '45J',
        '60': '60J', '60J': '60J',
        '90': '90J', '90J': '90J',
        '30JFM': '30J_FM', '30JFINDEMOIS': '30J_FM',
        'ACOMPTE30': 'ACOMPTE_30', 'ACOMPTE50': 'ACOMPTE_50',
    }
    if s in mapping:
        return mapping[s]
    try:
        n = int(float(s))
        if n in (30, 45, 60, 90):
            return f'{n}J'
        if n == 0:
            return 'COMPTANT'
    except Exception:
        pass
    return 'AUTRE'


def map_segment_from_secteur(secteur):
    s = safe_str(secteur).upper()
    if not s:
        return 'FLEXO'
    if 'HELIO' in s:
        return 'HELIO'
    if 'EXTRU' in s:
        return 'EXTRUSION'
    if 'INDUSTR' in s or 'AUTRE' in s:
        return 'AUTRE'
    if 'AGRO' in s or 'ALIMENT' in s:
        return 'FLEXO'
    return 'FLEXO'


def map_region_from_ville(ville):
    v = safe_str(ville).upper()
    v = unicodedata.normalize('NFKD', v).encode('ascii', 'ignore').decode('ascii')
    if not v:
        return ''

    ouest = [
        'ORAN', 'RELIZANE', 'GHLIZANE', 'MOSTAGANEM', 'MASCARA', 'SAIDA',
        'TIARET', 'TLEMCEN', 'SIDI BEL ABBES', 'SIDI BEL ABBES', 'SBA',
        'AIN TEMOUCHENT', 'TEMOUCHENT', 'MAGHNIA', 'ARZEW', 'ES SENIA',
        'BIR EL DJIR', 'ES-SENIA', 'GHAZAOUET', 'NEDROMA',
    ]
    centre = [
        'ALGER', 'ALGIERS', 'BLIDA', 'BOUMERDES', 'BOUMERDAS', 'TIPAZA', 'TIPASA',
        'TIZI', 'TIZI OUZOU', 'BOUIRA', 'MEDEA', 'MÉDÉA', 'AIN DEFLA',
        'CHLEF', 'ECHELIFF', 'EL KHEMIS', 'KOLEA', 'DAR EL BEIDA', 'ROUIBA',
        'REGHAIA', 'BORDJ EL KIFFAN', 'BAB EZZOUAR',
    ]
    est = [
        'CONSTANTINE', 'ANNABA', 'SETIF', 'SÉTIF', 'BEJAIA', 'BÉJAIA', 'BGA',
        'JIJEL', 'SKIKDA', 'GUELMA', 'OUED SOUF', "SOUK AHRAS", 'TEBESSA',
        'BATNA', 'KHENCHELA', 'OM BOUAGHI', "OUM EL BOUAGHI", 'MILA',
        'BORDJ BOU ARRERIDJ', 'BBA', 'EL EULMA',
    ]
    sud = [
        'OUARGLA', 'GHARDAIA', 'GHARDAÏA', 'BISKRA', 'LAGHOUAT', 'DJELFA',
        'BECHAR', 'BÉCHAR', 'ADRAR', 'TAMANRASSET', 'ILLIZI', 'TINDOUF',
        'EL OUED', 'OUED SOUF', 'TIMIMOUN', 'IN SALAH', 'DJANET',
        'HASSI MESSAOUD', 'TOUGGOURT', 'EL GOLEA',
    ]
    nord = [
        'AIN TAYA', 'ZERALDA', 'STAOUELI', 'DELLYS',
    ]

    def match_list(names):
        for n in names:
            if n in v or v in n:
                return True
        return False

    if match_list(ouest): return 'OUEST'
    if match_list(centre): return 'CENTRE'
    if match_list(est): return 'EST'
    if match_list(sud): return 'SUD'
    if match_list(nord): return 'NORD'
    if 'EXPORT' in v or 'ETRANGER' in v or 'FRANCE' in v or 'EUROPE' in v:
        return 'EXPORT'
    return ''


def resolve_commercial_user(commercial_raw):
    """
    Recherche uniquement parmi les utilisateurs ayant accès au module CRM.
    Si aucun n'est trouvé, retourne None (ne crée plus d'utilisateur automatiquement).
    """
    raw = safe_str(commercial_raw)
    if not raw:
        return None

    # Recherche uniquement parmi les utilisateurs CRM
    crm_users = get_crm_users_queryset()

    candidates = [
        raw, raw.lower(), raw.upper(), raw.replace(' ', ''),
        raw.replace(' ', '').lower(), raw.replace(' ', '_').lower(),
        re.sub(r'[^a-zA-Z0-9]', '', raw).lower(),
    ]
    compact = re.sub(r'[^a-zA-Z0-9]', '', raw).lower()
    if compact:
        candidates.append(compact)

    for cand in candidates:
        if not cand: continue
        user = crm_users.filter(
            Q(username__iexact=cand) |
            Q(username__icontains=cand) |
            Q(first_name__icontains=raw) |
            Q(last_name__icontains=raw)
        ).first()
        if user:
            return user

    return None


@login_required
def import_stock_view(request):
    # Charge uniquement les utilisateurs ayant accès au module CRM
    crm_users = get_crm_users_queryset()

    context = {
        'crm_users': crm_users
    }
    
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            import_type = request.POST.get('import_type')
            excel_file = request.FILES['excel_file']
            fs = FileSystemStorage()
            filename = fs.save(excel_file.name, excel_file)
            file_path = fs.path(filename)
            df = pd.read_excel(file_path).fillna('')

            rename_map = build_column_rename_map(df.columns)
            if rename_map and import_type == 'STOCK':
                df = df.rename(columns=rename_map)

            count = 0
            errors = 0
            details = []

            if import_type == 'STOCK':
                colonnes_reconnues = ', '.join(sorted(set(rename_map.values()))) if rename_map else None
                if colonnes_reconnues:
                    details.append(f"ℹ️ Colonnes Excel reconnues et mappées : {colonnes_reconnues}")
                else:
                    details.append("⚠️ Aucune colonne standard reconnue — vérifiez les en-têtes de votre fichier Excel.")

                for idx, row in df.iterrows():
                    try:
                        designation = str(row.get('Designation', 'Inconnu')).strip()
                        if designation in ('', '0', 'nan'):
                            designation = 'Inconnu'

                        raw_cat = str(row.get('Categorie', '')).strip()
                        if raw_cat in ('0', 'nan', 'None'):
                            raw_cat = ''

                        cat_code, cat_source = detect_material_category(raw_cat, designation)

                        Material.objects.update_or_create(
                            name=designation,
                            defaults={
                                'category': cat_code,
                                'quantity': float(row.get('Quantite', 0) or 0),
                                'unit': str(row.get('Unite', 'kg')).strip() or 'kg',
                                'min_threshold': float(row.get('Seuil_Min', 0) or 0),
                                'price_per_unit': float(row.get('Prix', 0) or 0)
                            }
                        )
                        count += 1
                        details.append(f"Ligne {idx+2}: ✅ {designation} importé en [{cat_code}] — détecté via {cat_source}")
                    except Exception as e:
                        errors += 1
                        details.append(f"Ligne {idx+2}: ❌ {str(e)}")

            elif import_type == 'CRM':
                df = auto_detect_crm_header(df)
                col_index = build_crm_column_index(df.columns)
                
                selected_commercial_id = request.POST.get('commercial_id')
                forced_commercial = None
                if selected_commercial_id:
                    # Vérifie que le commercial sélectionné a bien accès au CRM
                    forced_commercial = get_crm_users_queryset().filter(id=selected_commercial_id).first()

                if forced_commercial:
                    details.append(f"ℹ️ Tous les clients seront assignés au commercial : {forced_commercial.get_full_name() or forced_commercial.username}")
                else:
                    details.append(f"ℹ️ Détection auto du commercial selon le fichier Excel (uniquement parmi les utilisateurs ayant accès au CRM).")
                
                for idx, row in df.iterrows():
                    try:
                        code_client = crm_get(col_index, row, [
                            'ID Client', 'Id Client', 'id_client', 'Code Client', 'code_client', 'Code'
                        ])

                        nom = crm_get(col_index, row, [
                            'Nom Client', 'nom_client', 'Raison Sociale', 'raison_sociale',
                            'Nom', 'name', 'Raison_Sociale'
                        ])

                        if not nom:
                            details.append(f"Ligne {idx+2}: ⚠️ Nom Client vide — ignorée")
                            continue

                        secteur = crm_get(col_index, row, ['Secteur', 'sector', 'Activité', 'Activite'])
                        adresse = crm_get(col_index, row, ['Adresse', 'address', 'Adresse complete', 'Adresse complète'])
                        ville = crm_get(col_index, row, ['Ville', 'city', 'Wilaya'])
                        telephone = crm_get(col_index, row, [
                            'Téléphone', 'Telephone', 'Tel', 'phone', 'Tél', 'Teléphone'
                        ])
                        email = crm_get(col_index, row, ['Email', 'E-mail', 'mail', 'e_mail'])
                        statut_raw = crm_get(col_index, row, [
                            'Statut compte', 'Statut_compte', 'Statut', 'status', 'Etat', 'État'
                        ])
                        cond_paie_raw = crm_get(col_index, row, [
                            'Cond. paiement', 'Cond paiement', 'Conditions paiement',
                            'conditions_paiement', 'Paiement', 'Condition de paiement'
                        ])
                        lim_cred_raw = crm_get(col_index, row, [
                            'Limite crédit (DA)', 'Limite crédit', 'Limite credit (DA)',
                            'Limite credit', 'limite_credit', 'Crédit', 'Credit'
                        ])
                        observations = crm_get(col_index, row, [
                            'Observations', 'Notes', 'Remarques', 'Commentaire', 'notes'
                        ])
                        ice_nif = crm_get(col_index, row, [
                            'ICE', 'NIF', 'RC', 'ICE / NIF / RC', 'ice_nif', 'NIF/RC'
                        ])
                        date_1ere = crm_get(col_index, row, [
                            'Date 1ère cmd', 'Date 1ere cmd', 'Date premiere cmd',
                            'date_creation', 'Date entrée', 'Date entree'
                        ])
                        region_raw = crm_get(col_index, row, [
                            'Région', 'Region', 'Zone', 'region', 'wilaya_region'
                        ])

                        status = map_statut_compte(statut_raw)
                        conditions_paiement = map_conditions_paiement(cond_paie_raw)
                        segment = map_segment_from_secteur(secteur)

                        region = ''
                        if region_raw:
                            rr = region_raw.upper()
                            for code, label in [
                                ('NORD', 'NORD'), ('SUD', 'SUD'), ('EST', 'EST'),
                                ('OUEST', 'OUEST'), ('CENTRE', 'CENTRE'), ('EXPORT', 'EXPORT')
                            ]:
                                if code in rr:
                                    region = code
                                    break
                        if not region:
                            region = map_region_from_ville(ville)

                        try:
                            limite_credit = float(
                                lim_cred_raw.replace(' ', '').replace(',', '.').replace('DA', '')
                            ) if lim_cred_raw else 0.0
                        except (ValueError, TypeError):
                            limite_credit = 0.0

                        if forced_commercial:
                            commercial_user = forced_commercial
                        else:
                            commercial_raw = crm_get(col_index, row, ['Commercial', 'commercial', 'Vendeur'])
                            commercial_user = resolve_commercial_user(commercial_raw)

                        defaults = {
                            'name': nom[:200],
                            'sector': secteur[:100],
                            'address': adresse,
                            'city': ville[:100] if ville else 'Non renseignée',
                            'phone': telephone[:50] if telephone else '',
                            'email': email if ('@' in email) else '',
                            'status': status,
                            'segment': segment,
                            'region': region,
                            'limite_credit': limite_credit,
                            'notes': observations,
                            'conditions_paiement': conditions_paiement,
                        }
                        if ice_nif:
                            defaults['ice_nif'] = ice_nif[:100]
                        if commercial_user:
                            defaults['commercial'] = commercial_user

                        if date_1ere:
                            try:
                                d = pd.to_datetime(date_1ere, dayfirst=True, errors='coerce')
                                if pd.notna(d):
                                    defaults['date_creation'] = d.date()
                            except Exception:
                                pass

                        if code_client:
                            client_obj, created = Client.objects.update_or_create(
                                code_client=code_client[:50],
                                defaults=defaults
                            )
                            act = "créé" if created else "mis à jour"
                        else:
                            client_obj, created = Client.objects.update_or_create(
                                name=nom[:200],
                                defaults=defaults
                            )
                            act = "créé" if created else "mis à jour"

                        comm_label = (
                            commercial_user.get_full_name() or commercial_user.username
                            if commercial_user else '—'
                        )
                        reg_label = region or '—'
                        details.append(
                            f"Ligne {idx+2}: ✅ {nom} [{code_client or 'sans code'}] | "
                            f"{ville or '—'} | {reg_label} | Comm: {comm_label} | {act}"
                        )
                        count += 1
                    except Exception as e:
                        errors += 1
                        details.append(f"Ligne {idx+2}: ❌ Erreur : {str(e)}")

            elif import_type == 'SPECIAL_PROD':
                for idx, row in df.iterrows():
                    try:
                        try:
                            d_val = pd.to_datetime(row.get('Date')).date()
                        except Exception:
                            d_val = timezone.now().date()

                        produit_val = str(row.get('Produit', '')).strip()
                        if not produit_val or produit_val == '0':
                            details.append(f"Ligne {idx+2}: ⚠️ Produit vide — ignorée")
                            continue

                        client_obj = None
                        cli_name = str(row.get('Client', '')).strip()
                        if cli_name and cli_name not in ('0', ''):
                            client_obj, _ = Client.objects.get_or_create(
                                name=cli_name,
                                defaults={'city': 'Non renseigné', 'phone': '000'}
                            )

                        machine_obj = None
                        mac_name = str(row.get('Machine', '')).strip()
                        if mac_name and mac_name not in ('0', ''):
                            machine_obj = Machine.objects.filter(name__icontains=mac_name).first()
                            if not machine_obj:
                                machine_obj = Machine.objects.create(
                                    name=mac_name, type='IMP', status='STOP'
                                )
                                details.append(f"  → Machine '{mac_name}' créée automatiquement")

                        equipe_val = str(row.get('Equipe', 'A')).strip().upper()
                        if equipe_val not in ['A', 'B', 'C']:
                            equipe_val = 'A'

                        h_debut = parse_time_safe(row.get('H_Debut'))
                        h_fin = parse_time_safe(row.get('H_Fin'))

                        def safe_float(v, d=0):
                            try:
                                x = float(v)
                                return x if x == x else d
                            except Exception:
                                return d

                        entry = ProductionEntry.objects.create(
                            date=d_val, produit=produit_val,
                            support=str(row.get('Support', '')),
                            quantite_lancee=safe_float(row.get('Qte_Lancee')),
                            lot=str(row.get('Lot', '')),
                            laize=safe_float(row.get('Laize')),
                            client=client_obj, equipe=equipe_val,
                            machine=machine_obj,
                            heure_debut=h_debut, heure_fin=h_fin,
                            prod_ml=safe_float(row.get('Prod_ML')),
                            dechets_demarrage=safe_float(row.get('Dec_Demarrage')),
                            dechets_lisiere=safe_float(row.get('Dec_Lisiere')),
                            dechets_jonction=safe_float(row.get('Dec_Jonction')),
                            dechets_transport=safe_float(row.get('Dec_Transport')),
                            prod_kg=safe_float(row.get('Prod_KG')),
                            rebobinage_kg=safe_float(row.get('Rebobinage_KG'))
                        )
                        count += 1
                        details.append(f"Ligne {idx+2}: ✅ {produit_val} du {d_val} importé (ID={entry.id})")
                    except Exception as e:
                        errors += 1
                        details.append(f"Ligne {idx+2}: ❌ ERREUR — {str(e)}")

            elif import_type == 'CONSO':
                for idx, row in df.iterrows():
                    try:
                        try:
                            d_prod = pd.to_datetime(row.get('Date')).date()
                        except Exception:
                            d_prod = timezone.now().date()
                        raw_type = str(row.get('Type', 'FLEXO')).upper()
                        final_type = 'HELIO' if 'HELIO' in raw_type else 'FLEXO'
                        ConsommationEncre.objects.create(
                            date=d_prod, process_type=final_type,
                            support=row.get('Support', 'Inconnu'),
                            laize=float(row.get('Laize', 0)),
                            bobine_in=float(row.get('Bobine_In', 0)),
                            bobine_out=float(row.get('Bobine_Out', 0)),
                            metrage=float(row.get('Metrage', 0)),
                            encre_noir=float(row.get('Noir', 0)),
                            encre_magenta=float(row.get('Magenta', 0)),
                            encre_jaune=float(row.get('Jaune', 0)),
                            encre_cyan=float(row.get('Cyan', 0)),
                            encre_dore=float(row.get('Dore', 0)),
                            encre_silver=float(row.get('Silver', 0)),
                            encre_orange=float(row.get('Orange', 0)),
                            encre_blanc=float(row.get('Blanc', 0)),
                            encre_vernis=float(row.get('Vernis', 0)),
                            solvant_metoxyn=float(row.get('Metoxyn', 0)),
                            solvant_2080=float(row.get('2080', 0))
                        )
                        count += 1
                        details.append(f"Ligne {idx+2}: ✅ Conso {d_prod} importée")
                    except Exception as e:
                        errors += 1
                        details.append(f"Ligne {idx+2}: ❌ {str(e)}")

            elif import_type == 'TOOLS':
                for idx, row in df.iterrows():
                    try:
                        prod_ref = row.get('Ref_Produit')
                        if prod_ref and prod_ref != 0:
                            cli, _ = Client.objects.get_or_create(
                                name="Client Divers",
                                defaults={'city': 'Interne', 'phone': '000'}
                            )
                            product, _ = TechnicalProduct.objects.get_or_create(
                                ref_internal=prod_ref,
                                defaults={
                                    'name': f"Produit {prod_ref}",
                                    'client': cli, 'structure_type': 'MONO',
                                    'width_mm': 100
                                }
                            )
                            type_map = {'Cylindre': 'CYL', 'Cliche': 'CLICHE'}
                            type_val = row.get('Type', 'CYL')
                            if type_val == 0:
                                type_val = 'CYL'
                            Tooling.objects.update_or_create(
                                serial_number=row.get('Serial'),
                                defaults={
                                    'product': product,
                                    'tool_type': type_map.get(type_val, 'CYL'),
                                    'max_impressions': row.get('Tours_Max', 1000000),
                                    'current_impressions': row.get('Tours_Actuels', 0)
                                }
                            )
                            count += 1
                            details.append(f"Ligne {idx+2}: ✅ Outil {row.get('Serial')} importé")
                    except Exception as e:
                        errors += 1
                        details.append(f"Ligne {idx+2}: ❌ {str(e)}")

            elif import_type == 'PLANNING':
                for idx, row in df.iterrows():
                    try:
                        cli_name = row.get('Client')
                        if cli_name and cli_name != 0:
                            client, _ = Client.objects.get_or_create(
                                name=cli_name, defaults={'city': '?', 'phone': '?'}
                            )
                            prod_name = row.get('Produit')
                            product, _ = TechnicalProduct.objects.get_or_create(
                                ref_internal=f"REF-{str(prod_name)[:5]}",
                                defaults={
                                    'name': prod_name, 'client': client,
                                    'structure_type': 'MONO', 'width_mm': 500
                                }
                            )
                            mac_name = row.get('Machine')
                            machine = Machine.objects.filter(name__icontains=str(mac_name)).first()
                            try:
                                start_d = pd.to_datetime(row.get('Date_Debut'))
                            except Exception:
                                start_d = timezone.now()
                            end_d = start_d + timedelta(hours=4)
                            ProductionOrder.objects.update_or_create(
                                of_number=str(row.get('OF_Numero')),
                                defaults={
                                    'client': client, 'product': product,
                                    'machine': machine, 'start_time': start_d,
                                    'end_time': end_d,
                                    'quantity_planned': row.get('Qte_Prevue', 0),
                                    'status': 'PLANNED'
                                }
                            )
                            count += 1
                            details.append(f"Ligne {idx+2}: ✅ OF {row.get('OF_Numero')} importé")
                    except Exception as e:
                        errors += 1
                        details.append(f"Ligne {idx+2}: ❌ {str(e)}")

            context.update({
                'message': (
                    f'⚠️ {count} OK, {errors} erreurs.'
                    if errors else
                    f'✅ {count} lignes importées avec succès !'
                ),
                'success': count > 0,
                'details': details
            })
            try:
                os.remove(file_path)
            except Exception:
                pass

        except Exception as e:
            context.update({'message': f'❌ Erreur critique : {str(e)}', 'success': False})

    return render(request, 'stock/import_stock.html', context)


def download_template_special_prod(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Production Speciale"
    headers = [
        'Date', 'Produit', 'Support', 'Qte_Lancee', 'Lot', 'Laize',
        'Client', 'Equipe', 'Machine', 'H_Debut', 'H_Fin', 'Prod_ML',
        'Dec_Demarrage', 'Dec_Lisiere', 'Dec_Jonction', 'Dec_Transport',
        'Prod_KG', 'Rebobinage_KG'
    ]
    hf = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    hfill = PatternFill(start_color='0D47A1', end_color='0D47A1', fill_type='solid')
    tb = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(horizontal='center')
        cell.border = tb
    example = [
        '15/01/2025', 'Sac Lait 1L', 'PEBD 50μ', 500, 'LOT-001', 320,
        'Laiterie Atlas', 'A', 'IMP-01', '08:00', '16:30', 12000,
        5.2, 3.1, 1.5, 2.0, 485, 10
    ]
    ef = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    for col, val in enumerate(example, 1):
        cell = ws.cell(row=2, column=col, value=val)
        cell.fill = ef
        cell.border = tb
        cell.alignment = Alignment(horizontal='center')
    for col in ws.columns:
        ml = max((len(str(c.value)) for c in col if c.value), default=0)
        ws.column_dimensions[col[0].column_letter].width = ml + 4
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Template_Production_Speciale.xlsx"'
    wb.save(response)
    return response


def download_template_stock(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Matieres"
    headers = ['Designation', 'Categorie', 'Quantite', 'Unite', 'Seuil_Min', 'Prix']
    hf = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    hfill = PatternFill(start_color='0D47A1', end_color='0D47A1', fill_type='solid')
    tb = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(horizontal='center')
        cell.border = tb
    examples = [
        ['Encre Noir Flexo', 'Encre', 120, 'kg', 20, 850],
        ['Encre Magenta Helio', 'Encre', 45, 'kg', 15, 920],
        ['Film PEBD 50µ', 'Film', 3200, 'kg', 500, 210],
        ['Colle PU Bi-Composant', 'Colle', 60, 'kg', 10, 640],
        ['Solvant Metoxyn', 'Solvant', 200, 'l', 40, 180],
    ]
    ef = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    for row_idx, ex in enumerate(examples, 2):
        for col, val in enumerate(ex, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = ef
            cell.border = tb
            cell.alignment = Alignment(horizontal='center')
    for col in ws.columns:
        ml = max((len(str(c.value)) for c in col if c.value), default=0)
        ws.column_dimensions[col[0].column_letter].width = ml + 4
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Template_Stock_Matieres.xlsx"'
    wb.save(response)
    return response