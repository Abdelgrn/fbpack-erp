import io
from datetime import timedelta

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.messages import get_messages
from django.contrib import messages as django_messages
from django.db.models import Sum, Count, Avg, F, Q, FloatField, ExpressionWrapper
from django.db.models.functions import TruncMonth
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone

from ..models import (
    Client, CommandeClient, LigneCommandeClient, TechnicalProduct,
    FicheProductionJournaliere, Tooling, Material, Opportunite, Machine,
    ChatRoom, ChatMessage
)


SEUIL_USURE_OUTILLAGE = 80
SEUIL_TAUX_DECHET = 5.0


def _f(v):
    try:
        return float(v or 0)
    except Exception:
        return 0.0


def _purger_messages_kpi(request):
    """
    Supprime uniquement les bandeaux Django provenant des alertes KPI.
    Garde les autres messages normaux de l'ERP.
    """
    try:
        storage = get_messages(request)
        messages_a_garder = []

        for m in storage:
            txt = str(m)
            if '[KPI-ALERTE:' in txt:
                continue
            if 'Alertes Usine' in txt:
                continue
            if 'ChatMessage object' in txt:
                continue

            messages_a_garder.append({
                'level': m.level,
                'message': m.message,
                'extra_tags': m.extra_tags,
            })

        storage.used = True

        for m in messages_a_garder:
            django_messages.add_message(
                request,
                m['level'],
                m['message'],
                extra_tags=m['extra_tags']
            )
    except Exception:
        pass


def _ensure_alertes_usine_room():
    room, created = ChatRoom.objects.get_or_create(
        slug='alertes-usine',
        defaults={
            'name': 'Alertes Usine',
            'type': 'URGENCE',
            'icone': '🚨',
            'description': 'Alertes automatiques KPI : production, stock, outillage, retards',
            'est_actif': True,
        }
    )
    if not room.est_actif:
        room.est_actif = True
        room.save(update_fields=['est_actif'])
    return room


def _generer_alertes_kpi():
    alertes = []
    today = timezone.now().date()

    # ============================================================
    # 1. Outillage usé : clichés / cylindres
    # ============================================================
    try:
        outils = (
            Tooling.objects
            .select_related('product')
            .filter(max_impressions__gt=0)
            .annotate(
                pct=ExpressionWrapper(
                    F('current_impressions') * 100.0 / F('max_impressions'),
                    output_field=FloatField()
                )
            )
            .filter(pct__gte=SEUIL_USURE_OUTILLAGE)
            .order_by('-pct')[:20]
        )

        for t in outils:
            alertes.append({
                'key': f'OUTILLAGE-{t.id}',
                'niveau': 'critique' if t.pct >= 95 else 'warning',
                'type': 'OUTILLAGE',
                'icone': '⚙️',
                'titre': f"{t.get_tool_type_display()} {t.serial_number} usé à {t.pct:.0f}%",
                'detail': f"Produit : {t.product.name if t.product else 'Non lié'} · Tours {t.current_impressions}/{t.max_impressions}",
                'action': "Prévoir regravure / remplacement",
            })
    except Exception:
        pass

    # ============================================================
    # 2. Déchets élevés sur fiches production récentes
    # ============================================================
    try:
        fiches = (
            FicheProductionJournaliere.objects
            .filter(
                date_fabrication__gte=today - timedelta(days=30),
                total_poids_produit_kg__gt=0
            )
            .annotate(
                taux=ExpressionWrapper(
                    F('total_dechets_kg') * 100.0 / F('total_poids_produit_kg'),
                    output_field=FloatField()
                )
            )
            .filter(taux__gte=SEUIL_TAUX_DECHET)
            .order_by('-taux')[:20]
        )

        for f in fiches:
            alertes.append({
                'key': f'DECHETS-{f.id}',
                'niveau': 'critique' if f.taux >= SEUIL_TAUX_DECHET * 2 else 'warning',
                'type': 'DÉCHETS',
                'icone': '🗑️',
                'titre': f"Déchets élevés : {f.taux:.1f}% sur {f.numero_fiche}",
                'detail': f"{f.designation_produit or 'Produit non renseigné'} · {f.date_fabrication:%d/%m/%Y}",
                'action': "Analyser les causes de déchets",
            })
    except Exception:
        pass

    # ============================================================
    # 3. Commandes en retard
    # ============================================================
    try:
        retards = (
            CommandeClient.objects
            .select_related('client')
            .filter(date_livraison_prevue__lt=today)
            .exclude(statut__in=['LIVREE', 'ANNULEE'])
            .order_by('date_livraison_prevue')[:20]
        )

        for c in retards:
            jours = (today - c.date_livraison_prevue).days
            alertes.append({
                'key': f'COMMANDE-RETARD-{c.id}',
                'niveau': 'critique' if jours >= 7 else 'warning',
                'type': 'RETARD',
                'icone': '🚚',
                'titre': f"Commande {c.reference} en retard de {jours} jour(s)",
                'detail': f"{c.client.name} · prévue le {c.date_livraison_prevue:%d/%m/%Y} · {c.get_statut_display()}",
                'action': "Relancer production / logistique",
            })
    except Exception:
        pass

    # ============================================================
    # 4. Stock bas
    # ============================================================
    try:
        mats = (
            Material.objects
            .filter(min_threshold__gt=0, quantity__lte=F('min_threshold'))
            .order_by('quantity')[:20]
        )

        for m in mats:
            alertes.append({
                'key': f'STOCK-{m.id}',
                'niveau': 'warning',
                'type': 'STOCK',
                'icone': '📦',
                'titre': f"Stock bas : {m.name}",
                'detail': f"Stock actuel {m.quantity} {m.unit} · seuil {m.min_threshold}",
                'action': "Lancer demande d'achat / réapprovisionnement",
            })
    except Exception:
        pass

    alertes.sort(key=lambda a: 0 if a['niveau'] == 'critique' else 1)
    return alertes


def _push_alertes_to_chat(user, alertes):
    """
    Envoie les alertes KPI dans le salon Alertes Usine.
    Version silencieuse : aucun message Django affiché en haut du site.
    Anti-doublon par marqueur [KPI-ALERTE:key].
    """
    if not user or not user.is_authenticated:
        return 0

    room = _ensure_alertes_usine_room()
    sent = 0

    for a in alertes:
        marker = f"[KPI-ALERTE:{a['key']}]"

        exists = ChatMessage.objects.filter(
            room=room,
            contenu__icontains=marker
        ).exists()

        if exists:
            continue

        contenu = (
            f"{marker}\n"
            f"{a['icone']} {a['titre']}\n"
            f"📌 {a['detail']}\n"
            f"➡️ {a['action']}"
        )

        ChatMessage.objects.create(
            room=room,
            auteur=user,
            contenu=contenu,
            type_message='ALERT'
        )
        sent += 1

    return sent


def _calcul_oee_machines():
    resultats = {}

    fiches = (
        FicheProductionJournaliere.objects
        .select_related('machine')
        .filter(machine__isnull=False)
    )

    for f in fiches:
        mid = f.machine_id

        if mid not in resultats:
            resultats[mid] = {
                'machine_name': str(f.machine),
                'machine': str(f.machine),
                'nb_fiches': 0,
                'ouverture_min': 0.0,
                'fonctionnement_min': 0.0,
                'production': 0.0,
                'dechets': 0.0,
            }

        ouverture = _f(getattr(f, 'temps_ouverture_minutes', 0))
        if ouverture <= 0:
            ouverture = _f(f.t_preparation_min) + _f(f.t_fonctionnement_min) + _f(f.t_nettoyage_min)

        fonctionnement = _f(f.t_fonctionnement_min)
        production = _f(f.total_poids_produit_kg) or _f(getattr(f, 'total_prod_kg_calcul', 0))
        dechets = _f(f.total_dechets_kg)

        resultats[mid]['nb_fiches'] += 1
        resultats[mid]['ouverture_min'] += ouverture
        resultats[mid]['fonctionnement_min'] += fonctionnement
        resultats[mid]['production'] += production
        resultats[mid]['dechets'] += dechets

    rows = []

    for r in resultats.values():
        dispo = (r['fonctionnement_min'] / r['ouverture_min'] * 100) if r['ouverture_min'] else 0
        perf = 100 if r['fonctionnement_min'] > 0 else 0
        qual = (r['production'] / (r['production'] + r['dechets']) * 100) if (r['production'] + r['dechets']) > 0 else 0
        oee = dispo * perf * qual / 10000

        classe = 'world' if oee >= 85 else ('bon' if oee >= 65 else ('moyen' if oee >= 40 else 'faible'))

        rows.append({
            'machine_name': r['machine_name'],
            'machine': r['machine'],
            'nb_fiches': r['nb_fiches'],
            'heures_ouverture': round(r['ouverture_min'] / 60, 1),
            'heures_fonctionnement': round(r['fonctionnement_min'] / 60, 1),
            'heures_arret': round(max(r['ouverture_min'] - r['fonctionnement_min'], 0) / 60, 1),
            'dispo': round(dispo, 1),
            'perf': round(perf, 1),
            'qual': round(qual, 1),
            'oee': round(oee, 1),
            'classe': classe,
            'produit': r['production'],
            'dechets': r['dechets'],
        })

    rows.sort(key=lambda x: x['oee'], reverse=True)
    return rows


def _build_dashboard_context(request, push_to_chat=False):
    commandes_actives = CommandeClient.objects.exclude(statut='ANNULEE')

    ca_global = commandes_actives.aggregate(t=Sum('montant_total'))['t'] or 0
    nb_commandes_global = commandes_actives.count()

    clients_stats = {
        'total': Client.objects.count(),
        'actifs': Client.objects.filter(status='ACTIVE').count(),
        'vip': Client.objects.filter(status='VIP').count(),
    }

    produits_tech = {
        'total': TechnicalProduct.objects.filter(is_obsolete=False).count()
    }

    try:
        stock_total = {
            'total_items': Material.objects.count(),
            'total_qte': Material.objects.aggregate(s=Sum('quantity', output_field=FloatField()))['s'] or 0,
        }
    except Exception:
        stock_total = {'total_items': 0, 'total_qte': 0}

    fiches_stats = {
        'total_fiches': FicheProductionJournaliere.objects.count(),
        'dechet_moyen': FicheProductionJournaliere.objects.aggregate(m=Avg('total_dechets_kg'))['m'] or 0,
    }

    try:
        machines_stats = {
            'total': Machine.objects.count(),
            'actives': Machine.objects.filter(est_active=True).count() if hasattr(Machine, 'est_active') else Machine.objects.count(),
        }
    except Exception:
        machines_stats = {'total': 0, 'actives': 0}

    tooling_stats = {
        'total': Tooling.objects.count(),
        'cylindres': Tooling.objects.filter(tool_type='CYL').count(),
        'cliches': Tooling.objects.filter(tool_type='CLICHE').count(),
    }

    top_clients_commande = (
        Client.objects
        .annotate(
            nb_cmd=Count('commandes_client'),
            ca=Sum(
                'commandes_client__montant_total',
                filter=Q(commandes_client__statut__in=['CONFIRMEE', 'LIVREE', 'PRETE', 'EN_PRODUCTION'])
            )
        )
        .filter(nb_cmd__gt=0)
        .order_by('-ca')[:10]
    )

    top_produits_vendus = (
        LigneCommandeClient.objects
        .select_related('produit')
        .values('produit__name', 'produit__ref_internal')
        .annotate(
            qte_totale=Sum('quantite', output_field=FloatField()),
            nb_lignes=Count('id')
        )
        .order_by('-qte_totale')[:10]
    )

    produits_dechets = (
        FicheProductionJournaliere.objects
        .exclude(designation_produit='')
        .values('designation_produit')
        .annotate(
            nb_fiches=Count('id'),
            total_dechets=Sum('total_dechets_kg', output_field=FloatField()),
            duree_moyenne=Avg('t_fonctionnement_min')
        )
        .order_by('-total_dechets')[:10]
    )

    top_tooling = (
        Tooling.objects
        .select_related('product')
        .filter(max_impressions__gt=0)
        .annotate(
            pct_calc=ExpressionWrapper(
                F('current_impressions') * 100.0 / F('max_impressions'),
                output_field=FloatField()
            )
        )
        .order_by('-current_impressions')[:10]
    )

    oee_machines = _calcul_oee_machines()
    oee_global = round(sum([x['oee'] for x in oee_machines]) / len(oee_machines), 1) if oee_machines else 0

    opportunites_stats = {
        'total': Opportunite.objects.count(),
        'gagnees': Opportunite.objects.filter(status='GAGNE').count(),
        'perdues': Opportunite.objects.filter(status='PERDU').count(),
    }

    alertes_list = _generer_alertes_kpi()
    alertes_critiques = sum(1 for a in alertes_list if a['niveau'] == 'critique')

    if push_to_chat:
        _push_alertes_to_chat(request.user, alertes_list)

    # Graphiques
    now = timezone.now().date()
    start_12 = now - timedelta(days=365)

    ca_mois = (
        commandes_actives
        .filter(date_commande__gte=start_12)
        .annotate(mois=TruncMonth('date_commande'))
        .values('mois')
        .annotate(total=Sum('montant_total'))
        .order_by('mois')
    )

    dechets_mois = (
        FicheProductionJournaliere.objects
        .filter(date_fabrication__gte=start_12)
        .annotate(mois=TruncMonth('date_fabrication'))
        .values('mois')
        .annotate(total=Sum('total_dechets_kg'))
        .order_by('mois')
    )

    segments = (
        commandes_actives
        .values('client__segment')
        .annotate(total=Sum('montant_total'))
        .order_by('-total')
    )

    graph_data = {
        'ca_labels': [x['mois'].strftime('%m/%Y') for x in ca_mois if x['mois']],
        'ca_values': [_f(x['total']) for x in ca_mois],
        'dechets_labels': [x['mois'].strftime('%m/%Y') for x in dechets_mois if x['mois']],
        'dechets_values': [_f(x['total']) for x in dechets_mois],
        'segment_labels': [x['client__segment'] or 'Autre' for x in segments],
        'segment_values': [_f(x['total']) for x in segments],
        'oee_labels': [x['machine_name'] for x in oee_machines],
        'oee_values': [x['oee'] for x in oee_machines],
    }

    context = {
        'ca_global': ca_global,
        'nb_commandes_global': nb_commandes_global,
        'clients_stats': clients_stats,
        'produits_tech': produits_tech,
        'stock_total': stock_total,
        'fiches_stats': fiches_stats,
        'machines_stats': machines_stats,
        'tooling_stats': tooling_stats,
        'top_clients_commande': top_clients_commande,
        'top_produits_vendus': top_produits_vendus,
        'produits_dechets': produits_dechets,
        'top_tooling': top_tooling,
        'oee_machines': oee_machines,
        'oee_list': oee_machines,
        'oee_global': oee_global,
        'opportunites_stats': opportunites_stats,
        'alertes_list': alertes_list,
        'alertes': alertes_list,
        'alertes_critiques': alertes_critiques,
        'graph_data': graph_data,
        'now': timezone.now(),
    }

    return context


@login_required
def kpi_dashboard(request):
    _purger_messages_kpi(request)
    context = _build_dashboard_context(request, push_to_chat=True)
    return render(request, 'kpi/dashboard.html', context)


@login_required
def kpi_client(request, client_id):
    client = get_object_or_404(Client, id=client_id)
    return redirect(f"{reverse('kpi_dashboard')}?client={client.id}&periode=all")


@login_required
def kpi_oee(request):
    _purger_messages_kpi(request)
    context = _build_dashboard_context(request, push_to_chat=False)
    return render(request, 'kpi/oee.html', context)


@login_required
def kpi_export_excel(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return HttpResponse("La librairie openpyxl est requise : pip install openpyxl", status=500)

    ctx = _build_dashboard_context(request, push_to_chat=False)
    wb = Workbook()

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E293B")

    def feuille(ws, titre, entetes, lignes):
        ws.title = titre[:31]
        ws.append(entetes)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal='center')
        for l in lignes:
            ws.append(l)
        for col in ws.columns:
            largeur = max(len(str(c.value or '')) for c in col) + 3
            ws.column_dimensions[col[0].column_letter].width = min(largeur, 50)

    feuille(wb.active, "Synthèse", ["Indicateur", "Valeur"], [
        ["CA Global", ctx['ca_global']],
        ["Commandes actives", ctx['nb_commandes_global']],
        ["Clients total", ctx['clients_stats']['total']],
        ["Clients actifs", ctx['clients_stats']['actifs']],
        ["Produits techniques", ctx['produits_tech']['total']],
        ["Fiches production", ctx['fiches_stats']['total_fiches']],
        ["Outillages", ctx['tooling_stats']['total']],
        ["Alertes critiques", ctx['alertes_critiques']],
    ])

    feuille(
        wb.create_sheet(),
        "Top Clients",
        ["Client", "Code", "Commandes", "CA"],
        [[c.name, c.code_client, c.nb_cmd, c.ca] for c in ctx['top_clients_commande']]
    )

    feuille(
        wb.create_sheet(),
        "Produits Vendus",
        ["Produit", "Réf.", "Quantité", "Lignes"],
        [[p['produit__name'], p['produit__ref_internal'], p['qte_totale'], p['nb_lignes']] for p in ctx['top_produits_vendus']]
    )

    feuille(
        wb.create_sheet(),
        "Déchets",
        ["Produit", "Fiches", "Déchets", "Durée moyenne"],
        [[p['designation_produit'], p['nb_fiches'], p['total_dechets'], p['duree_moyenne']] for p in ctx['produits_dechets']]
    )

    feuille(
        wb.create_sheet(),
        "OEE Machines",
        ["Machine", "Fiches", "Ouverture h", "Fonctionnement h", "Dispo", "Perf", "Qualité", "OEE"],
        [[o['machine_name'], o['nb_fiches'], o['heures_ouverture'], o['heures_fonctionnement'], o['dispo'], o['perf'], o['qual'], o['oee']] for o in ctx['oee_machines']]
    )

    feuille(
        wb.create_sheet(),
        "Alertes",
        ["Niveau", "Type", "Titre", "Détail", "Action"],
        [[a['niveau'], a['type'], a['titre'], a['detail'], a['action']] for a in ctx['alertes_list']]
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nom = f"KPI_FBPACK_{timezone.now():%Y%m%d}.xlsx"
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{nom}"'
    return response


@login_required
def kpi_rapport_print(request):
    _purger_messages_kpi(request)
    context = _build_dashboard_context(request, push_to_chat=False)
    return render(request, 'kpi/rapport_print.html', context)


@login_required
def fiche_technique_print(request, product_id):
    produit = get_object_or_404(TechnicalProduct.objects.select_related('client'), id=product_id)
    couleurs = produit.couleurs_flexo.all()
    outillages = produit.tooling_set.all()

    for t in outillages:
        t.pct_affiche = t.wear_percent

    return render(request, 'kpi/fiche_technique_print.html', {
        'produit': produit,
        'couleurs': couleurs,
        'outillages': outillages,
        'now': timezone.now(),
    })


@login_required
def kpi_fiche_technique(request, product_id):
    return fiche_technique_print(request, product_id)